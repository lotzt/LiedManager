"""Startet die Oberfläche ohne Bildschirm (offscreen), prüft die wichtigsten Abläufe und speichert Screenshots.

Abgedeckt: Schriften (Arial/Georgia, nur mit echter Bildschirm-Plattform), Start, Liedsuche, Ansagen/Pausentimer,
Einstellungen (inkl. Datenbank-Wechsel), Seitenleiste mit den Reitern Kamera/Bilder, Kamera-Status sowie
die Bilder-Funktion (Ablage, Vorschau, Beamer-Galerie).

Aufruf:  python tools/smoke_test.py            (Screenshots landen in build/screenshots)
"""
import os
import sqlite3
import sys
import tempfile
import time
import traceback
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SHOTS = ROOT / "build" / "screenshots"


def make_test_db(folder: Path, name: str = "LiedManager.db", songs=None) -> None:
    con = sqlite3.connect(folder / name)
    con.execute("CREATE TABLE Data (nummer TEXT, alias TEXT, titel TEXT, lyrics TEXT, "
                "audio1 BLOB, audio2 BLOB, bild1 BLOB, bild2 BLOB, bild3 BLOB, bild4 BLOB)")
    con.executemany("INSERT INTO Data (nummer, alias, titel, lyrics) VALUES (?,?,?,?)", songs or [
        ("p001", None, "Großer Gott, wir loben dich", "Großer Gott, wir loben dich,\nHerr, wir preisen deine Stärke."),
        ("p002", None, "Befiehl du deine Wege", "Befiehl du deine Wege\nund was dein Herze kränkt"),
    ])
    con.commit()
    con.close()


def make_test_images(folder: Path) -> None:
    """Drei Testbilder in unterschiedlichen Formaten (Querformat, Hochformat, groß) plus eine Textdatei."""
    from PyQt6.QtGui import QColor, QImage, QPainter

    folder.mkdir()
    for name, (w, h), color in (("bild1.png", (1600, 900), "#C08A53"),
                                ("bild2.png", (900, 1400), "#5A7D9A"),
                                ("bild10.png", (5000, 3000), "#7A9A5A")):
        image = QImage(w, h, QImage.Format.Format_RGB32)
        image.fill(QColor(color))
        painter = QPainter(image)
        painter.setPen(QColor("white"))
        font = painter.font()
        font.setPixelSize(h // 6)
        painter.setFont(font)
        painter.drawText(image.rect(), 0x84, name)   # AlignCenter
        painter.end()
        assert image.save(str(folder / name)), f"{name} konnte nicht geschrieben werden"
    (folder / "notizen.txt").write_text("kein Bild", encoding="utf-8")


def test_updater_logic(tmp: Path) -> None:
    """Update-Logik ohne Oberfläche und ohne Internet: Versionsvergleich, Dateiwahl, Download (lokaler Server)."""
    import hashlib
    import http.server
    import threading

    from services import updater as up

    assert up.parse_version("v2.10.1") == (2, 10, 1) and up.parse_version("kaputt") is None
    assert up.is_newer("v2.0.10", "2.0.9") and up.is_newer("v3.0.0", "2.9.9")
    assert not up.is_newer("v2.0.0", "2.0.0") and not up.is_newer("v1.9.9", "2.0.0") and not up.is_newer("x", "2.0.0")
    assert up.asset_suffix("win32", "AMD64") == "-Windows-Setup.exe"
    assert up.asset_suffix("darwin", "arm64") == "-macOS-arm64.dmg"
    assert up.asset_suffix("darwin", "x86_64") == "-macOS-x86_64.dmg"
    assert up.asset_suffix("linux", "x86_64") is None
    assert up.safe_name("../../evil.exe") == "evil.exe"

    payload = os.urandom(700_000)
    digest = hashlib.sha256(payload).hexdigest()
    served = tmp / "served"
    served.mkdir()
    (served / "LiedManager-9.9.9-Windows-Setup.exe").write_bytes(payload)

    class Quiet(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=str(served), **kwargs)

        def log_message(self, *args):
            pass

    class Server(http.server.ThreadingHTTPServer):
        def handle_error(self, request, client_address):   # Abbrüche (Cancel-Test) sind hier gewollt
            pass

    server = Server(("127.0.0.1", 0), Quiet)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        def release(tag, names=("LiedManager-9.9.9-Windows-Setup.exe", "LiedManager-9.9.9-macOS-arm64.dmg")):
            return {"tag_name": tag, "body": "Neu: Updater", "html_url": "https://example.invalid/r",
                    "assets": [{"name": n, "size": len(payload), "digest": f"sha256:{digest}",
                                "browser_download_url": f"{base}/{n}"} for n in names]}

        info = up.check("o/r", "2.0.0", fetch=lambda repo: release("v9.9.9"), suffix="-Windows-Setup.exe")
        assert info.version == "9.9.9" and info.sha256 == digest and info.size == len(payload)
        assert up.check("o/r", "9.9.9", fetch=lambda repo: release("v9.9.9"), suffix="-Windows-Setup.exe") is None
        for bad, expected in ((lambda: up.check("", "2.0.0"), "nicht eingerichtet"),
                              (lambda: up.check("o/r", "2.0.0", fetch=lambda r: release("v9.9.9", ("x.zip",)),
                                                suffix="-Windows-Setup.exe"), "fehlt noch"),
                              (lambda: up.check("o/r", "2.0.0", fetch=lambda r: release("nightly"),
                                                suffix="-Windows-Setup.exe"), "Unbekannte Versionsnummer")):
            try:
                bad()
            except up.UpdateError as exc:
                assert expected in str(exc), str(exc)
            else:
                raise AssertionError(f"UpdateError „{expected}“ erwartet")

        dest = tmp / "updates"
        seen = []
        path = up.download(info, dest, progress=lambda read, total: seen.append((read, total)))
        assert path.read_bytes() == payload and not list(dest.glob("*.part")), "Download unvollständig"
        assert seen and seen[-1][0] == len(payload)

        from dataclasses import replace
        for broken, expected in ((replace(info, sha256="0" * 64), "Prüfsumme"), (replace(info, sha256=""), "Prüfsumme"),
                                 (replace(info, size=len(payload) + 1), "unvollständig")):
            try:
                up.download(broken, dest)
            except up.UpdateError as exc:
                assert expected in str(exc), str(exc)
            else:
                raise AssertionError(f"Download-Fehler „{expected}“ erwartet")
        try:
            up.download(info, dest, cancelled=lambda: True)
        except up.UpdateCancelled:
            pass
        else:
            raise AssertionError("Abbruch wurde ignoriert")
        assert not list(dest.glob("*.part")), "Teil-Download wurde nicht aufgeräumt"
        up.cleanup(dest)
        assert not list(dest.iterdir()), "cleanup hat nicht aufgeräumt"
    finally:
        server.shutdown()


class FakeSyncServer:
    """Kleiner Schein-Server mit denselben Regeln wie server/sync.php (für Tests ohne PHP und MySQL)."""

    TOKEN = "test-token"
    LIST_LIMIT = 2            # absichtlich klein, damit das Blättern in mehreren Seiten mitgetestet wird

    def __init__(self):
        import http.server
        import json
        import threading

        self.rows = {}        # key -> dict(ts, dev, hash, deleted, seq, data)
        self.counter = 0
        self.pushes = 0
        self.parts, self.part_sizes = {}, []
        self.reject = set()   # Schlüssel, die der Server mit „zu groß“ ablehnt
        fake = self

        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def _send(self, code, payload=None):
                body = json.dumps(payload).encode("utf-8") if payload is not None else b""
                self.send_response(code)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def do_POST(self):
                data = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or b"{}")
                if data.get("token") != fake.TOKEN:
                    return self._send(401)
                action = data.get("action")
                if action == "list":
                    items = sorted((r for r in fake.rows.values() if r["seq"] > int(data.get("since", 0))),
                                   key=lambda r: r["seq"])
                    more = len(items) > fake.LIST_LIMIT
                    items = items[:fake.LIST_LIMIT]
                    return self._send(200, {"format": 1, "max": fake.counter, "more": more, "items": [
                        {"k": r["key"], "ts": r["ts"], "dev": r["dev"], "deleted": r["deleted"], "seq": r["seq"]}
                        for r in items]})
                if action == "push_part":
                    fake.parts[(data["uid"], data["idx"])] = data["data"]
                    fake.part_sizes.append(len(data["data"]))
                    return self._send(200, {"ok": True})
                if action == "get_part":
                    part = fake.parts.get((data["uid"], data["idx"]))
                    return self._send(200, {"data": part}) if part is not None else self._send(404, {"error": "fehlt"})
                if action == "get":
                    row = fake.rows.get(data.get("key"))
                    if row and row.get("uid"):
                        return self._send(200, {"row": dict(row, k=row["key"], data=None, parts=row["parts"])})
                    return self._send(200, {"row": dict(row, k=row["key"]) if row else None})
                if action == "push":
                    fake.pushes += 1
                    key = data["key"]
                    if key in fake.reject:
                        return self._send(413, {"error": "zu groß"})
                    old = fake.rows.get(key)
                    if old and (old["ts"], old["dev"]) >= (data["ts"], data["dev"]):
                        return self._send(200, {"accepted": False})
                    if data.get("parts") and any((data["uid"], i) not in fake.parts for i in range(data["parts"])):
                        return self._send(400, {"error": "Teile des Liedes fehlen."})
                    fake.counter += 1
                    fake.rows[key] = {"key": key, "ts": data["ts"], "dev": data["dev"], "hash": data["hash"],
                                      "deleted": 1 if data.get("deleted") else 0, "seq": fake.counter,
                                      "data": None if data.get("deleted") else data.get("data"),
                                      "uid": data.get("uid"), "parts": data.get("parts")}
                    return self._send(200, {"accepted": True, "seq": fake.counter})
                return self._send(400, {"error": "Unbekannte Aktion."})

        self.httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}/sync.php"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def stored_bytes(self) -> int:
        return (sum(len(r["data"] or "") for r in self.rows.values())
                + sum(len(p) for p in self.parts.values()))

    def stop(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def test_server_sync(tmp: Path) -> None:
    """Abgleich über den Server mit simulierten Geräten (ohne Oberfläche, gegen einen Schein-Server)."""
    import shutil

    from services import server_sync as ss

    def make_db(path, rows):
        con = sqlite3.connect(path)
        con.execute("CREATE TABLE Data (nummer TEXT, alias TEXT, titel TEXT, lyrics TEXT, "
                    "audio1 BLOB, audio2 BLOB, bild1 BLOB, bild2 BLOB, bild3 BLOB, bild4 BLOB)")
        con.executemany("INSERT INTO Data (nummer, alias, titel, lyrics, audio1) VALUES (?,?,?,?,?)", rows)
        con.commit()
        con.close()

    def dump(path):
        con = sqlite3.connect(path)
        rows = con.execute("SELECT nummer, alias, titel, lyrics, audio1 FROM Data ORDER BY nummer").fetchall()
        con.close()
        return rows

    def edit(path, sql, args=()):
        con = sqlite3.connect(path)
        con.execute(sql, args)
        con.commit()
        con.close()

    assert ss.url_problem("https://technik.example.de/sync/sync.php") == ""
    assert ss.url_problem("http://example.de/sync.php") and ss.url_problem("") and ss.url_problem("example.de")
    assert ss.url_problem("http://127.0.0.1:8000/sync.php") == "", "lokaler Test-Server ohne HTTPS muss erlaubt sein"

    work = tmp / "abgleich"
    work.mkdir(parents=True)
    server = FakeSyncServer()
    try:
        big = os.urandom(1_500_000)
        base = [("p001", None, "Eins", "Text 1", None), ("p002", None, "Zwei", "Text 2", None),
                ("p004", None, "Vier mit Audio", "Text 4", big)]
        a_db, b_db, c_db = work / "A.db", work / "B.db", work / "C.db"
        make_db(a_db, base)
        shutil.copyfile(a_db, b_db)
        make_db(c_db, base)
        clock = {"t": 1_000_000}

        def tick():
            clock["t"] += 1000
            return clock["t"]

        def device(db, name, now_ms=tick, token=FakeSyncServer.TOKEN):
            return ss.ServerSync(db, server.url, token, work / f"{name}_state.db", now_ms=now_ms)

        dev_a, dev_b, dev_c = device(a_db, "A"), device(b_db, "B"), device(c_db, "C")

        r = dev_a.run_once()
        assert r.ok and r.baseline and r.sent == 0, r
        assert dev_b.run_once().baseline
        assert not server.rows, "erster Lauf darf nichts senden (sonst landet die ganze Datenbank auf dem Server)"

        edit(a_db, "UPDATE Data SET titel = 'Eins neu' WHERE nummer = 'p001'")
        edit(a_db, "INSERT INTO Data (nummer, titel, lyrics, audio1) VALUES ('p003', 'Drei', 'Text 3', ?)", (os.urandom(40_000),))
        edit(a_db, "DELETE FROM Data WHERE nummer = 'p002'")
        r = dev_a.run_once()
        assert r.ok and r.sent == 3, r
        assert len(server.rows) == 3 and server.stored_bytes() < 120_000, \
            "es dürfen nur die geänderten Lieder verschickt werden, nicht das unveränderte Audio von p004"

        app = sqlite3.connect(b_db)                        # so hält das Programm die Datenbank offen
        before = app.execute("PRAGMA data_version").fetchone()[0]
        r = dev_b.run_once()
        assert r.ok and r.received == 3, r
        assert dump(a_db) == dump(b_db), "Geräte A und B müssen denselben Stand haben"
        assert app.execute("PRAGMA data_version").fetchone()[0] != before, "das laufende Programm muss die Änderung bemerken"
        app.close()

        for dev in (dev_a, dev_b):
            r = dev.run_once()
            assert r.ok and r.sent == 0 and r.received == 0, f"nichts darf doppelt gesendet/übernommen werden: {r}"

        # Medien, die als Text-Typ mit ungültigem UTF-8 gespeichert sind (kam in echten Datenbanken vor)
        edit(a_db, "INSERT INTO Data (nummer, titel, bild1) VALUES ('p005', 'Mit Bild', CAST(x'FFD8FFE000104A46' AS TEXT))")
        r = dev_a.run_once()
        assert r.ok and r.sent == 1, r
        r = dev_b.run_once()
        assert r.ok and r.received == 1, r

        def raw_image(path):
            con = sqlite3.connect(path)
            value = con.execute("SELECT CAST(bild1 AS BLOB) FROM Data WHERE nummer = 'p005'").fetchone()[0]
            con.close()
            return value
        assert raw_image(a_db) == raw_image(b_db) == bytes.fromhex("FFD8FFE000104A46"), "Bilddaten müssen unverändert ankommen"
        for dev in (dev_a, dev_b):
            r = dev.run_once()
            assert r.ok and r.sent == 0 and r.received == 0, f"nichts darf doppelt gesendet/übernommen werden: {r}"

        edit(a_db, "UPDATE Data SET titel = 'von A' WHERE nummer = 'p001'")
        edit(b_db, "UPDATE Data SET titel = 'von B' WHERE nummer = 'p001'")
        dev_a.run_once()
        dev_b.run_once()
        dev_a.run_once()
        assert dump(a_db) == dump(b_db), "bei Streit um dasselbe Lied müssen am Ende alle gleich sein"
        assert dump(a_db)[0][2] == "von B", "die jüngere Änderung muss gewinnen"

        edit(a_db, "UPDATE Data SET lyrics = 'neu' WHERE nummer = 'p004'")
        dev_a.run_once()
        lock = sqlite3.connect(b_db)
        lock.execute("BEGIN EXCLUSIVE")
        old_timeout, ss.LOCK_TIMEOUT = ss.LOCK_TIMEOUT, 0.3
        try:
            r = dev_b.run_once()
        finally:
            ss.LOCK_TIMEOUT = old_timeout
            lock.rollback()
            lock.close()
        assert not r.ok, "eine gesperrte Datenbank muss als Fehler gemeldet werden, ohne Absturz"
        r = dev_b.run_once()
        assert r.ok and r.received == 1 and dump(a_db) == dump(b_db), r

        r = dev_c.run_once()                               # neues Gerät mit alter Datenbank holt alles nach (mehrere Seiten)
        assert r.ok and r.baseline and r.received >= 3 and dump(c_db) == dump(a_db), r

        # Gerät mit nachgehender Uhr: seine Änderung muss trotzdem gewinnen, weil es den Stand der anderen schon kennt
        slow = device(c_db, "C", now_ms=lambda: 5)
        edit(c_db, "UPDATE Data SET titel = 'von C (Uhr geht nach)' WHERE nummer = 'p001'")
        r = slow.run_once()
        assert r.ok and r.sent == 1, r
        dev_a.run_once()
        dev_b.run_once()
        assert dump(a_db) == dump(b_db) == dump(c_db), "alle Geräte müssen gleich sein"
        assert dump(a_db)[0][2] == "von C (Uhr geht nach)", "Änderung bei nachgehender Uhr darf nicht verloren gehen"

        # „Alles neu hochladen“: Der Stand von Gerät C wird für alle verbindlich (inkl. Löschungen)
        edit(c_db, "DELETE FROM Data WHERE nummer = 'p003'")
        edit(c_db, "UPDATE Data SET titel = 'Eins aus der neuen Datenbank' WHERE nummer = 'p001'")
        edit(c_db, "INSERT INTO Data (nummer, titel, lyrics) VALUES ('p009', 'Neun', 'Text 9')")
        steps = []
        forced = ss.ServerSync(c_db, server.url, FakeSyncServer.TOKEN, work / "C_state.db", now_ms=tick, force_upload=True,
                               on_progress=lambda done, total: steps.append((done, total)))
        r = forced.run_once()
        assert steps and steps[-1][0] == steps[-1][1] > 0 and steps == sorted(steps), f"Fortschritt unplausibel: {steps}"
        assert r.ok and r.forced and r.sent >= 4 and "Stand für alle Geräte gesetzt" in r.text(), r
        dev_a.run_once()
        dev_b.run_once()
        assert dump(a_db) == dump(b_db) == dump(c_db), "nach „Alles neu hochladen“ müssen alle Geräte gleich sein"
        numbers = [row[0] for row in dump(a_db)]
        assert "p003" not in numbers and "p009" in numbers, "Löschung und neues Lied müssen bei allen ankommen"
        assert dump(a_db)[0][2] == "Eins aus der neuen Datenbank"
        r = dev_c.run_once()
        assert r.ok and r.sent == 0 and r.received == 0, f"danach darf nichts mehr zu tun sein: {r}"
        leer = work / "leer.db"
        make_db(leer, [])
        r = ss.ServerSync(leer, server.url, FakeSyncServer.TOKEN, work / "L_state.db", force_upload=True).run_once()
        assert not r.ok and "keine Lieder" in r.error, "eine leere Datenbank darf den Server nicht leeren"
        assert len(numbers) == len([k for k, row in server.rows.items() if not row["deleted"]]), "Server und Geräte müssen übereinstimmen"

        # Der Server lehnt ein Lied ab (zu groß): kein Absturz, keine ständige Wiederholung
        server.reject.add("p004")
        edit(a_db, "UPDATE Data SET lyrics = 'zu groß' WHERE nummer = 'p004'")
        r = dev_a.run_once()
        assert r.ok and r.failed == 1 and "abgelehnt" in r.text(), r
        pushes = server.pushes
        r = dev_a.run_once()
        assert r.ok and r.failed == 1 and server.pushes == pushes, "abgelehntes Lied darf nicht dauernd neu hochgeladen werden"
        server.reject.clear()
        retry = ss.ServerSync(a_db, server.url, FakeSyncServer.TOKEN, work / "A_state.db", now_ms=tick, retry_failed=True)
        r = retry.run_once()
        assert r.ok and r.failed == 0 and r.sent == 1, f"manueller Abgleich muss abgelehnte Lieder sofort erneut senden: {r}"
        dev_a.run_once()                                   # Zwischenstand: das „zu groß“-Lied erneut versuchen lassen
        # Großes Lied wird in Teilen hochgeladen und beim Empfänger wieder zusammengesetzt
        old_part = ss.PART_BYTES
        ss.PART_BYTES = 100_000
        try:
            edit(a_db, "UPDATE Data SET audio1 = ?, lyrics = 'mit großem Audio' WHERE nummer = 'p004'", (os.urandom(900_000),))
            server.part_sizes.clear()
            r = dev_a.run_once()
            assert r.ok and r.sent == 1, r
            assert len(server.part_sizes) >= 10 and max(server.part_sizes) <= 100_000, server.part_sizes
            assert server.rows["p004"]["data"] is None and server.rows["p004"]["parts"] == len(server.part_sizes)
            r = dev_b.run_once()
            assert r.ok and r.received == 1, r
            assert dump(a_db) == dump(b_db), "großes Lied muss nach dem Zusammensetzen identisch ankommen"
        finally:
            ss.PART_BYTES = old_part
        dev_a.run_once()

        # CCLI-Nummern: Gerät A bekommt eine Spalte und eine Nummer; Gerät B (ohne Spalte) übernimmt sie, auch das Löschen.
        # Lieder ohne CCLI-Nummer behalten ihren Fingerabdruck – es wird nichts unnötig neu gesendet.
        assert dev_a.run_once().sent == 0 and dev_b.run_once().sent == 0
        edit(a_db, "ALTER TABLE Data ADD COLUMN CCLI TEXT")
        r = dev_a.run_once()
        assert r.ok and r.sent == 0, f"leere CCLI-Spalte darf nichts auslösen: {r}"
        edit(a_db, "UPDATE Data SET CCLI = ' 123 456 ' WHERE nummer = 'p001'")
        r = dev_a.run_once()
        assert r.ok and r.sent == 1, r
        r = dev_b.run_once()
        assert r.ok and r.received == 1, r
        con = sqlite3.connect(b_db)
        assert con.execute("SELECT CCLI FROM Data WHERE nummer = 'p001'").fetchone() == ("123456",), "CCLI nicht bei B angekommen"
        con.close()
        assert dev_b.run_once().sent == 0, "Übernahme darf nicht zurückgespielt werden"
        edit(a_db, "UPDATE Data SET CCLI = NULL WHERE nummer = 'p001'")
        assert dev_a.run_once().sent == 1
        dev_b.run_once()
        con = sqlite3.connect(b_db)
        assert con.execute("SELECT CCLI FROM Data WHERE nummer = 'p001'").fetchone() == (None,), "Löschen kam nicht an"
        con.close()
        assert dev_b.run_once().sent == 0

        # Falsches Token, geleerter Server
        r = device(a_db, "A", token="falsch").run_once()
        assert not r.ok and "Token" in r.error, r
        server.rows.clear()
        server.counter = 0
        r = dev_b.run_once()
        assert r.ok, f"ein geleerter Server darf keinen Fehler auslösen: {r}"

        server.stop()
        r = dev_a.run_once()
        assert not r.ok and "erreichbar" in r.error, "fehlender Server muss gemeldet werden"
    finally:
        try:
            server.stop()
        except Exception:
            pass


def test_remote_relay(tmp: Path) -> None:
    """Website-Steuerung ohne Oberfläche: der Client im Programm gegen eine nachgebaute remote.php."""
    import base64
    import http.server
    import json
    import socket
    import threading

    from services import remote_relay as rr
    from services.remote_common import RemoteError

    token = "geheimes-token-0123456789"

    # Meldung der Website bei fehlenden Zugangsdaten: die genaue Ursache soll im Programm ankommen
    import io
    import urllib.error
    body = io.BytesIO(json.dumps({"error": "Noch nicht eingerichtet: PIN ist zu kurz (mindestens 8 Zeichen, eingetragen: 6)."}).encode())
    err = urllib.error.HTTPError("https://example.org/remote.php", 503, "Service Unavailable", {}, body)
    assert "PIN ist zu kurz" in rr.RemoteApi._http_message(err), "503 muss die Ursache der Website zeigen"
    empty = urllib.error.HTTPError("https://example.org/remote.php", 503, "Service Unavailable", {}, io.BytesIO(b""))
    assert "noch nicht eingerichtet" in rr.RemoteApi._http_message(empty), "503 ohne Text braucht eine Ersatzmeldung"

    class Site:
        """Stand der „Website“: Aufträge, Ergebnisse, letzter Zustand."""

        def __init__(self, token_value=token, fmt=rr.FORMAT):
            self.token, self.format = token_value, fmt
            self.jobs, self.results, self.state, self.bye = [], {}, None, 0
            self.lock = threading.Lock()
            site = self

            class Handler(http.server.BaseHTTPRequestHandler):
                def log_message(self, *args):
                    pass

                def do_POST(self):
                    data = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                    if data.get("token") != site.token:
                        self.send_response(401)
                        self.send_header("Content-Length", "0")
                        self.end_headers()
                        return
                    action, reply = data.get("action"), {"ok": True}
                    if action == "poll":
                        end = time.monotonic() + 0.3
                        while time.monotonic() < end and not site.jobs:
                            time.sleep(0.02)
                        with site.lock:
                            jobs, site.jobs = site.jobs, []
                        reply = {"format": site.format, "jobs": jobs}
                    elif action == "state":
                        site.state = json.loads(data["data"])
                    elif action == "result":
                        site.results[data["id"]] = (data["http"], data["ctype"], base64.b64decode(data["body"]))
                    elif action == "bye":
                        site.bye += 1
                    payload = json.dumps(reply).encode()
                    self.send_response(200)
                    self.send_header("Content-Type", "application/json")
                    self.send_header("Content-Length", str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)

            self.httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Handler)
            self.url = f"http://127.0.0.1:{self.httpd.server_port}/remote.php"
            threading.Thread(target=self.httpd.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True).start()

        def close(self):
            self.httpd.shutdown()
            self.httpd.server_close()

    class Fake:
        def __init__(self):
            self.calls = []

        def state(self):
            return {"version": "test"}

        def action(self, name, args):
            self.calls.append((name, args))
            if name == "boom":
                raise RemoteError("Das geht gerade nicht.", 409)
            return {"ok": True, "echo": args}

        def thumbnail(self, index):
            return b"\xff\xd8\xff" if index == 0 else None

        def upload(self, filename, data):
            self.calls.append(("upload", filename, len(data)))
            return {"added": 1, "failed": 0}

    def wait_for(condition, seconds=8.0):
        end = time.monotonic() + seconds
        while time.monotonic() < end:
            if condition():
                return True
            time.sleep(0.02)
        return False

    saved = (rr.STATE_EVERY, rr.MIN_CYCLE)
    rr.STATE_EVERY, rr.MIN_CYCLE = 0.1, 0.05
    backend = Fake()
    sites, clients = [], []
    try:
        # Eingaben prüfen: ohne HTTPS und mit kurzem Token startet nichts
        assert "https://" in rr.RelayClient("http://example.com/remote.php", token, backend).start()
        assert "mindestens" in rr.RelayClient("https://example.com/remote.php", "kurz", backend).start()

        # falsches Token wird verständlich gemeldet
        site = Site()
        sites.append(site)
        texts = []
        wrong = rr.RelayClient(site.url, "x" * 20, backend, on_status=texts.append)
        clients.append(wrong)
        assert wrong.start() == ""
        assert wait_for(lambda: any("Token falsch" in text for text in texts)), texts
        wrong.stop()

        # Programm und remote.php passen nicht zusammen
        old = Site(fmt=99)
        sites.append(old)
        texts = []
        mismatch = rr.RelayClient(old.url, token, backend, on_status=texts.append)
        clients.append(mismatch)
        mismatch.start()
        assert wait_for(lambda: any("passen nicht zusammen" in text for text in texts)), texts
        mismatch.stop()

        # Website nicht erreichbar
        probe = socket.socket()
        probe.bind(("127.0.0.1", 0))
        closed = probe.getsockname()[1]
        probe.close()
        texts = []
        lost = rr.RelayClient(f"http://127.0.0.1:{closed}/remote.php", token, backend, on_status=texts.append)
        clients.append(lost)
        lost.start()
        assert wait_for(lambda: any("Keine Verbindung" in text for text in texts)), texts
        lost.stop()

        # Normalbetrieb
        site = Site()
        sites.append(site)
        texts = []
        client = rr.RelayClient(site.url, token, backend, on_status=texts.append)
        clients.append(client)
        assert client.start() == "" and client.running
        assert wait_for(lambda: client.connected and site.state is not None), texts
        assert site.state == {"version": "test"} and any("Verbunden" in text and site.url in text for text in texts)

        ids = iter(range(1, 100))

        def run(kind, name, body=""):
            job_id = next(ids)
            site.jobs.append({"id": job_id, "kind": kind, "name": name, "body": body})
            assert wait_for(lambda: job_id in site.results), f"Auftrag {kind}/{name} wurde nicht beantwortet"
            return site.results[job_id]

        status, ctype, data = run("action", "escape", '{"a": 1}')
        assert status == 200 and ctype.startswith("application/json") and json.loads(data)["echo"] == {"a": 1}
        assert backend.calls[-1] == ("escape", {"a": 1})
        status, _, data = run("action", "boom", "{}")
        assert status == 409 and json.loads(data)["error"] == "Das geht gerade nicht."
        assert run("action", "x", "kein json")[0] == 400, "ungültiges JSON muss abgelehnt werden"
        assert json.loads(run("action", "x", "[1]")[2])["echo"] == {}, "Liste statt Objekt zählt als leer"
        status, ctype, data = run("thumb", "0")
        assert (status, ctype, data) == (200, "image/jpeg", b"\xff\xd8\xff")
        assert run("thumb", "5")[0] == 404 and run("thumb", "abc")[0] == 400
        picture = base64.b64encode(b"BILD").decode()
        status, _, data = run("upload", "foto.JPG", picture)
        assert status == 200 and json.loads(data) == {"added": 1, "failed": 0}
        assert backend.calls[-1] == ("upload", "foto.JPG", 4)
        assert run("upload", "virus.exe", picture)[0] == 415
        assert run("upload", "leer.png", "")[0] == 400 and run("upload", "kaputt.png", "!!!")[0] == 400
        assert run("seltsam", "x")[0] == 400

        # frischer Zustand nach jedem Ergebnis, und das Beenden meldet das Programm bei der Website ab
        client.stop()
        assert not client.running and site.bye == 1, "beim Beenden muss die Website informiert werden"
    finally:
        rr.STATE_EVERY, rr.MIN_CYCLE = saved
        for item in clients:
            item.stop()
        for item in sites:
            item.close()


def test_remote_site() -> None:
    """server/remote.php stammt aus Vorlage + Webseite, ist aktuell und enthält keine echten Zugangsdaten."""
    import importlib.util
    import re

    from ui.web_remote_ui import design_values

    spec = importlib.util.spec_from_file_location("build_remote_site", ROOT / "tools" / "build_remote_site.py")
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)

    assert builder.design_values() == design_values(), "Design-Werte des Generators weichen vom Programm ab"
    committed = (ROOT / "server" / "remote.php").read_text(encoding="utf-8")
    assert committed == builder.render(), "server/remote.php ist veraltet: python tools/build_remote_site.py ausführen"
    for placeholder in ("'HIER-LANGES-ZUFAELLIGES-TOKEN'", "'HIER-PASSWORT-FUER-DEN-BROWSER'", "'HIER-DB-PASSWORT'"):
        assert placeholder in committed, f"{placeholder} fehlt – in das Repository gehören nie echte Zugangsdaten"
    assert not re.search(r"__PAGE_[A-Z]+__", committed), "Platzhalter der Vorlage nicht ersetzt"
    for tag in ("LM_PAGE_HTML", "LM_PAGE_CSS", "LM_PAGE_JS"):
        assert len(re.findall(rf"^{tag};?$|<<<'{tag}'", committed, re.M)) == 2, f"{tag}: Anfang/Ende stimmt nicht"
    assert "X-LM-Token" in committed and "?asset=js" in committed and "?asset=css" in committed
    assert "die du in der Datei remote.php festgelegt hast" in committed, "Anmeldetext der Website-Ausführung fehlt"


def make_test_pptx(path: Path) -> None:
    """Kleine Pausenfolien-Vorlage wie die echte: Ansagen, Bibelvers, Countdown „60“, „Ende!“ – mit Bild und Farben."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    from PyQt6.QtGui import QColor, QImage
    picture = path.with_suffix(".png")
    image = QImage(64, 64, QImage.Format.Format_RGB32)
    image.fill(QColor("#335577"))
    image.save(str(picture))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(18288000), Emu(10287000)
    layout = prs.slide_layouts[6]

    def box(slide, name, text, x, y, w, h, size, color="323232"):
        shape = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        shape.name = name
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(color)
        return shape

    first = prs.slides.add_slide(layout)
    first.shapes.add_picture(str(picture), 0, 0, Emu(18288000), Emu(10287000))
    box(first, "TextBox 4", "Ansagen", 1747110, 309926, 5392916, 2206181, 139)
    box(first, "AnsagenTextfeld", "Platzhalter", 1282360, 3133156, 10661768, 1342034, 34)
    second = prs.slides.add_slide(layout)
    box(second, "TextBox 4", "Bibelvers", 1747110, 309926, 6482490, 2206181, 139)
    verse = box(second, "TextBox 6", "Ehre sei Gott in der Höhe", 1758832, 3439268, 9366367, 7705827, 39)
    verse.text_frame.add_paragraph().add_run().text = "\t\tLukas 2,14"
    third = prs.slides.add_slide(layout)
    third.shapes.add_picture(str(picture), Emu(1028700), Emu(773485), Emu(8729105), Emu(8740030))
    box(third, "TextBox 5", "60", 3676956, 2399575, 3432592, 5340821, 337, "404040")
    last = prs.slides.add_slide(layout)
    box(last, "TextBox 5", "Ende!", 345002, 2517076, 10096500, 5252848, 275, "404040")
    prs.save(str(path))


def test_pausenfolie(tmp, window, app) -> None:
    """Pausenfolie ohne PowerPoint: Vorlage übernehmen, Ablauf nach Restzeit, Anzeige auf dem Beamer, Live-Ansagen."""
    from services import pausenfolie
    from ui import pause_window
    from services.settings import app_settings as _settings
    source = tmp / "pausentimer_test.pptx"
    make_test_pptx(source)
    design = pausenfolie.import_pptx(source, "test")
    assert set(design["scenes"]) == set(pausenfolie.SCENES), "es fehlen Ansichten"
    roles = {layer.get("role") for scene in design["scenes"].values() for layer in scene["layers"]
             if layer["kind"] == "text"}
    assert {"announcements", "verse", "number", "end", "static"} <= roles, f"Text-Rollen fehlen: {roles}"
    assert design["verse_text"] == "Ehre sei Gott in der Höhe" and design["verse_reference"] == "Lukas 2,14"
    assert design["timing"]["countdown_seconds"] == 60
    number = pausenfolie.text_layer(design, "countdown", "number")
    assert number and "center_in" in number, "die Zahl gehört mittig in den Kreis"
    assert [pausenfolie.phase(s, design) for s in (400, 300, 61, 60, 0.4, 0)] == [
        ("announcements", None), ("verse", None), ("verse", None), ("countdown", 60), ("countdown", 1), ("end", None)]
    previews = pause_window.preview_images("test", (320, 180), ["Erste Ansage", "Zweite Ansage"])
    assert [name for name, _img in previews] == ["Ansagen", "Bibelvers", "Countdown", "Ende"]
    assert all(not img.isNull() and img.width() == 320 for _name, img in previews), "Vorschau wird nicht gezeichnet"
    for name, img in previews:
        img.save(str(SHOTS / f"pausenfolie_{name}.png"))
    assert pause_window._pulse_values(0.0) == (1.0, 1.0) and pause_window._pulse_values(0.5) == (1.05, 1.0), \
        "Pulsieren darf die Zahl nicht durchsichtig machen"
    assert design["pulse"] is False, "wie in PowerPoint: Zahlen wechseln ohne Effekt"

    # Ablauf mit echtem Fenster (Musik gibt es in der Testvorlage nicht)
    _settings().setValue("presentation_season", "test")
    _settings().setValue(pause_window.KEY_MODE, "liedmanager")
    timer = window.pausentimer
    from PyQt6.QtCore import QTime
    import datetime as _dt
    end = (_dt.datetime.now() + _dt.timedelta(minutes=20)).time()
    timer.start(["Hallo"], QTime(end.hour, end.minute))
    assert timer.running and timer._active is timer._native, "Pausentimer ohne PowerPoint startet nicht"
    native = timer._native
    app.processEvents()
    assert native._window is not None and native._window.isVisible() and native._window.scene_name == "announcements"
    assert native._window.texts.get("announcements") == "Hallo", "Ansagen müssen schon beim Start auf dem Beamer stehen"
    window.pausentimer_panel.text_edit.setPlainText("Neue Ansage")
    assert native._texts["announcements"] == "Neue Ansage", "Ansagen müssen live auf den Beamer"
    native._finish = _dt.datetime.now() + _dt.timedelta(seconds=30.5)
    native._tick()
    assert native._window.scene_name == "countdown" and native._window.number == 31
    native._window.grab().save(str(SHOTS / "pausenfolie_beamer.png"))
    native._finish = _dt.datetime.now() - _dt.timedelta(seconds=1)
    native._tick()
    assert native._window.scene_name == "end"
    timer.stop()
    assert not timer.running and native._window is None, "Pausentimer stoppt nicht"
    _settings().setValue("presentation_season", "frühling")
    window.pausentimer_panel.text_edit.setPlainText("")


def test_history_and_problem_report(tmp) -> None:
    """Liedhistorie (Datei lesen/suchen/zählen) und automatische Problemmeldung (Entdoppeln, Warteschlange, Datenschutz)."""
    import tempfile
    from pathlib import Path

    from services import history, problem_report

    file = Path(tempfile.mkdtemp(dir=str(tmp))) / "g.txt"
    file.write_text("Sonntag, 18.10.2026\r\n12 – Großer Gott\r\n\r\nSonntag, 25.10.2026\r\n12 – Großer Gott\r\n99\r\n", encoding="utf-8")
    days = history.read(file)
    assert [d[0] for d in days] == ["Sonntag, 25.10.2026", "Sonntag, 18.10.2026"], days
    assert history.top_songs(days)[0] == ("12", "Großer Gott", 2)
    assert len(history.filter_days(days, "groß")) == 2 and history.filter_days(days, "zzz") == []
    assert "anna" not in problem_report.scrub("/Users/anna/x token=abc")
    assert "abc" not in problem_report.scrub("token=abc")
    assert problem_report.fingerprint("E 5", "File a, line 1") == problem_report.fingerprint("E 7", "File a, line 9")


def test_xair() -> None:
    """X AIR: OSC-Pakete und Steuerung gegen ein Schein-Pult (UDP auf dem eigenen Rechner)."""
    import socket
    import threading
    import time

    from services import xair

    assert xair.osc_decode(xair.osc_encode("/ch/01/mix/fader", 0.75)) == ("/ch/01/mix/fader", [0.75])
    assert xair.osc_decode(xair.osc_encode("/ch/03/mix/on", 0)) == ("/ch/03/mix/on", [0])
    assert xair.osc_decode(xair.osc_encode("/ch/02/config/name", "Pastor")) == ("/ch/02/config/name", ["Pastor"])
    assert abs(xair.fader_db(0.75) - 0.0) < 1e-6 and xair.fader_db(0) == float("-inf")
    server = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    server.bind(("127.0.0.1", 0))
    server.settimeout(0.3)
    received, stop = [], threading.Event()
    values = {"/ch/02/mix/fader": 0.6, "/ch/02/mix/on": 1, "/ch/02/config/name": "Pastor"}

    def fake_mixer():
        while not stop.is_set():
            try:
                data, addr = server.recvfrom(4096)
            except socket.timeout:
                continue
            decoded = xair.osc_decode(data)
            received.append(decoded)
            if decoded and not decoded[1] and decoded[0] in values:
                server.sendto(xair.osc_encode(decoded[0], values[decoded[0]]), addr)

    thread = threading.Thread(target=fake_mixer, daemon=True)
    thread.start()
    mixer = xair.XAirMixer("127.0.0.1", server.getsockname()[1])
    mixer.start()
    try:
        deadline = time.time() + 5
        while time.time() < deadline and not (mixer.connected and mixer.state()["channels"][1]["name"]):
            time.sleep(0.05)
        channel = mixer.state()["channels"][1]
        assert mixer.connected and channel["name"] == "Pastor" and abs(channel["level"] - 0.6) < 1e-3 and channel["mute"] is False
        mixer.set_level(2, 0.4)
        mixer.set_mute(2, True)
        time.sleep(0.4)
        assert any(r and r[1] and r[0] == "/ch/02/mix/fader" and abs(r[1][0] - 0.4) < 1e-3 for r in received), "Fader-Befehl fehlt"
        assert ("/ch/02/mix/on", [0]) in received, "Mute-Befehl fehlt"
        received.clear()
        mixer.mute_all()
        time.sleep(0.4)
        assert all((f"/ch/{n:02d}/mix/on", [0]) in received for n in range(1, 9)), "Alle-stumm fehlt"
        try:
            mixer.set_level(9, 0.5)
            raise AssertionError("Kanal 9 darf es nicht geben")
        except ValueError:
            pass
    finally:
        mixer.stop()
        stop.set()
        server.close()


def test_new_features(tmp: Path) -> None:
    """Lied mit fehlerhaftem Text behält seine Medien, Aufnahmen-Upload mit Fortsetzen, vorgemerkte Bilder abholen."""
    import base64

    from database import DatabaseManager
    from services import recording_upload as up
    from services import remote_relay as relay

    # 1) Lied p12: ungültiges UTF-8 im Text darf nichts löschen und nicht abstürzen
    folder = tmp / "p12"
    folder.mkdir()
    make_test_db(folder, songs=[("p012", None, "Test", "Text")])
    con = sqlite3.connect(folder / "LiedManager.db")
    con.execute("UPDATE Data SET audio1 = ?, bild1 = ?, alias = ? WHERE nummer = 'p012'",
                (b"AUDIO" * 10, b"BILD" * 10, b"\xff\xfe kaputt"))
    con.commit()
    con.close()
    db = DatabaseManager(str(folder / "LiedManager.db"))
    row = db.get_song_full_data(1)
    assert row is not None and row[4] == "Text", "Lied mit fehlerhaftem Text muss lesbar sein"
    db.close()
    con = sqlite3.connect(folder / "LiedManager.db")
    assert con.execute("SELECT audio1, bild1 FROM Data").fetchone() == (b"AUDIO" * 10, b"BILD" * 10), \
        "Beim Lesen dürfen Audio und Bilder nie gelöscht werden"
    con.close()

    # 2) Aufnahme-Upload: setzt dort fort, wo die Website aufgehört hat
    sent = []

    class FakeUploadApi:
        def __init__(self, *args):
            pass

        def call(self, action, **params):
            sent.append((action, params.get("offset"), params.get("last")))
            return {"size": 3_000_000} if action == "rec_size" else {"ok": True, "name": params["name"]}

    original = up.RemoteApi
    up.RemoteApi = FakeUploadApi
    try:
        wav = tmp / "2026.10.25 - Test.wav"
        wav.write_bytes(b"0" * 5_000_000)
        assert up.upload_file("https://example.org/remote.php", "t" * 20, str(wav)) == wav.name
    finally:
        up.RemoteApi = original
    assert sent[0][0] == "rec_size" and sent[1][1] == 3_000_000 and sent[-1][2] is True, sent

    # 3) Von zu Hause hochgeladene Bilder (Zwischenspeicher auf der Website) werden abgeholt und bestätigt
    items = [{"id": 1, "name": "a.png", "body": base64.b64encode(b"x").decode()},
             {"id": 2, "name": "b.txt", "body": base64.b64encode(b"y").decode()}]
    calls, taken = [], []

    class FakeInboxApi:
        def __init__(self, *args):
            pass

        def call(self, action, **params):
            calls.append(action)
            return {"item": items.pop(0) if items else None} if action == "inbox_get" else {"ok": True}

    class FakeBackend:
        def upload(self, name, data):
            taken.append(name)

    original = relay.RemoteApi
    relay.RemoteApi = FakeInboxApi
    try:
        client = relay.RelayClient("https://example.org/remote.php", "t" * 20, FakeBackend())
        client._inbox_busy = True
        client._drain_inbox()
    finally:
        relay.RemoteApi = original
    assert taken == ["a.png"] and calls.count("inbox_done") == 2 and not client._inbox_busy, (taken, calls)


def main() -> int:
    import faulthandler

    # Eigene Sicherung: hängt der Test (z. B. wegen eines Dialogs, den niemand bestätigen kann), steht im Log,
    # wo er steht – und er endet nach SMOKE_TEST_TIMEOUT Sekunden, statt ewig zu laufen.
    faulthandler.dump_traceback_later(int(os.environ.get("SMOKE_TEST_TIMEOUT", "240")), exit=True)

    tmp = Path(tempfile.mkdtemp(prefix="liedmanager_smoke_"))
    os.environ["LIEDMANAGER_SETTINGS_DIR"] = str(tmp)   # der Test darf nie die echten Einstellungen lesen oder ändern
    make_test_db(tmp)
    os.chdir(tmp)  # die Testdatenbank wird zuerst im Arbeitsverzeichnis gesucht
    SHOTS.mkdir(parents=True, exist_ok=True)
    from services import songlog as _songlog
    _songlog_dir = tmp / "Gespielte Lieder"                 # die echte Liederliste des Benutzers bleibt unberührt
    _songlog_dir.mkdir()
    _songlog.songlog_dir = lambda: _songlog_dir

    from PyQt6.QtCore import QAbstractAnimation, QSettings
    from PyQt6.QtWidgets import QApplication

    QSettings.setPath(QSettings.Format.NativeFormat, QSettings.Scope.UserScope, str(tmp))
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(tmp))

    from ui.main_window import MainWindow
    from ui.settings_dialog import SettingsDialog
    from ui.theme import apply_theme
    from ui.tooltip import install_tooltips

    app = QApplication(sys.argv)
    apply_theme(app)
    install_tooltips(app)

    # Einen modalen Dialog kann im Test niemand bestätigen: lieber sofort scheitern als für immer hängen
    from PyQt6.QtWidgets import QDialog

    def no_modal_dialog(self, *args, **kwargs):
        raise AssertionError(f"Unerwarteter modaler Dialog im Test: {type(self).__name__}")

    QDialog.exec = no_modal_dialog

    # ---- Schriften: auf Windows und macOS müssen Arial und Georgia vorhanden sein -------------
    from PyQt6.QtGui import QFontDatabase, QFontInfo

    from ui import theme
    installed = set(QFontDatabase.families())
    if app.platformName() == "offscreen":
        # Die Offscreen-Plattform kennt die Systemschriften nicht (auf Windows-Runnern ist die Liste leer); in der
        # echten App sind Arial und Georgia vorhanden und werden nicht mitgeliefert (Microsoft-Lizenz).
        print("Schriftenprüfung übersprungen (Offscreen-Plattform kennt die Systemschriften nicht)")
    elif sys.platform in ("win32", "darwin"):
        for family in theme.SAFE_FAMILIES:
            assert family in installed, f"Schrift {family} fehlt auf diesem System"
    body = QFontInfo(theme.body_font()).family()
    print(f"Schriften: Text = {theme.resolve_family(theme.BODY_FAMILIES) or body}, "
          f"Überschriften = {theme.resolve_family(theme.SERIF_FAMILIES) or '(Systemstandard)'}")

    def settle():
        for _ in range(5):
            app.processEvents()

    # ---- Start, Suche, Ansagen, Einstellungen -------------------------------------------------
    window = MainWindow()
    window.show()
    settle()
    window.grab().save(str(SHOTS / "01_start.png"))

    assert window.db.available, "Testdatenbank nicht gefunden"
    window.search_input.setText("1")
    window._on_search_return_pressed()
    settle()
    assert window.current_song_summary is not None, "Lied 1 wurde nicht gefunden"
    window.grab().save(str(SHOTS / "02_lied_gewaehlt.png"))

    window.pausentimer_panel.toggle()
    settle()
    window.grab().save(str(SHOTS / "03_ansagen_offen.png"))

    window.search_input.setText("gott")
    window._on_search_return_pressed()
    settle()
    window.grab().save(str(SHOTS / "04_suche.png"))

    window.handle_escape()
    settle()
    assert window.search_input.text() == "", "Esc hat die Suche nicht geleert"

    dialog = SettingsDialog(window)
    dialog.show()
    for index, name in enumerate(("allgemein", "pausentimer", "lieder", "audio", "kamera", "remote")):
        dialog.tab_group.button(index).click()
        settle()
        dialog.grab().save(str(SHOTS / f"05_einstellungen_{name}.png"))
    dialog.tab_group.button(1).click()
    from PyQt6.QtWidgets import QButtonGroup as _QBG
    subtabs = [g for g in dialog.pages.widget(1).findChildren(_QBG) if g.buttons()]
    assert subtabs and len(subtabs[0].buttons()) == 5, "Pausentimer: Vorlage / Vorschau / Ablauf / Musik / Schrift"
    for index in range(5):
        subtabs[0].button(index).click()
        settle()
        dialog.grab().save(str(SHOTS / f"05_einstellungen_pausentimer_{index + 1}.png"))
    dialog.tab_group.button(4).click()
    settle()
    assert abs(dialog.cam1_combo.width() - dialog.cam2_combo.width()) <= 1, "Kamera 1 und 2 müssen gleich breit sein"
    dialog.tab_group.button(3).click()
    settle()
    assert abs(dialog.audio_device_combo.width() - dialog.mic_device_combo.width()) <= 1, "Ausgabe und Eingabe gleich breit"
    dialog.close()

    import time

    from PyQt6.QtCore import Qt

    from ui.side_panel import TAB_CAMERA, TAB_IMAGES
    camera_panel = window.camera_panel
    assert window.side_panel.current_tab() == TAB_CAMERA, "Seitenleiste muss mit dem Reiter Kamera starten"
    assert camera_panel.isVisible(), "Kamera-Panel ist nicht sichtbar"
    assert all(slot.focusPolicy() == Qt.FocusPolicy.ClickFocus for slot in camera_panel.slots), \
        "Kamera-Rahmen dürfen nur per Klick den Fokus bekommen"
    assert not any(slot.hasFocus() for slot in camera_panel.slots), "Kamera-Rahmen ist ohne Klick angewählt"
    assert not hasattr(camera_panel, "status"), "die Statuszeile „Live“ unter den Kameras ist entfernt"
    for slot in camera_panel.slots:
        assert slot.height() == window.side_panel._slot_h, \
            f"Kamerafeld hat nicht die feste Höhe: {slot.width()}x{slot.height()}"
        assert not hasattr(slot, "live_dot"), "der rote Kamera-Punkt ist entfernt"
    # Kameras enden an der Pausentimer-Kopfzeile; „Keine Bilder“ steht mittig in deren Höhe -> Abstände oben/unten gleich
    from PyQt6.QtCore import QPoint as _P
    _win = window
    _cam_bottom = camera_panel.mapTo(_win, _P(0, camera_panel.height())).y()
    _head = window.pausentimer_panel.title
    _head_top = _head.mapTo(_win, _P(0, 0)).y()
    assert _cam_bottom == _head_top, f"Kameras enden nicht an der Pausentimer-Kopfzeile ({_cam_bottom} / {_head_top})"
    _box = window.side_panel.images_count_box
    assert _box.mapTo(_win, _P(0, 0)).y() == _head_top and _box.height() == _head.height(), \
        "Anzahl-Zeile liegt nicht genau auf Höhe der Pausentimer-Kopfzeile"
    slot = camera_panel.slots[0]
    slot.set_state("live")                      # Zustand „Bild läuft“ nachstellen (offscreen gibt es keine Kamera)
    settle()
    from PyQt6.QtCore import QPoint
    from PyQt6.QtTest import QTest
    assert not hasattr(camera_panel, "sync_live_dots"), "der rote Kamera-Punkt ist entfernt"
    window.side_panel.show_tab(TAB_IMAGES)
    window.side_panel.show_tab(TAB_CAMERA)

    # Doppelklick auf ein laufendes Kamerabild = auf den Beamer
    fired = []
    slot.fullscreen_toggled.connect(fired.append)
    QTest.mouseDClick(slot, Qt.MouseButton.LeftButton)
    assert fired == [0], "Doppelklick löst das Beamer-Vollbild nicht aus"
    camera_panel.close_all_fullscreen()         # sonst setzt das spätere Schließen das Bild wieder auf „läuft“
    slot.set_state("empty")
    QTest.mouseDClick(slot, Qt.MouseButton.LeftButton)
    assert fired == [0], "Doppelklick ohne Kamera darf nichts auslösen"
    slot.fullscreen_toggled.disconnect(fired.append)
    settle()

    # ---- Bilder-Liste: sanftes Scrollen, kein Springen beim Klicken ---------------------------
    from PyQt6.QtGui import QPixmap
    from PyQt6.QtWidgets import QListWidgetItem

    from ui.images_panel import _ThumbList
    from ui.widgets import SmoothListWidget

    def SmoothListWidget_wheel():
        return SmoothListWidget.wheelEvent
    assert _ThumbList.wheelEvent is not SmoothListWidget_wheel(), "Bilderliste scrollt nicht selbst: das Rad gehört dem äußeren Feld"
    thumbs = _ThumbList()
    thumbs.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)     # Testaufbau: Rad-Logik der Basisklasse prüfen
    thumbs.wheelEvent = lambda event, _f=SmoothListWidget_wheel(): _f(thumbs, event)
    for _ in range(20):
        item = QListWidgetItem()
        item.setData(Qt.ItemDataRole.UserRole, QPixmap(10, 10))
        thumbs.addItem(item)
    thumbs.resize(240, 300)
    thumbs.show()
    settle()
    bar = thumbs.verticalScrollBar()
    assert bar.maximum() > 200, "Testliste muss scrollbar sein"
    bar.setValue(0)
    partly_hidden = thumbs.visualItemRect(thumbs.item(2)).center().y() + 40    # Bild 3 ragt unten aus der Liste heraus
    QTest.mouseClick(thumbs.viewport(), Qt.MouseButton.LeftButton, pos=QPoint(100, min(partly_hidden, 290)))
    QTest.mouseDClick(thumbs.viewport(), Qt.MouseButton.LeftButton, pos=QPoint(100, min(partly_hidden, 290)))
    settle()
    assert bar.value() == 0, f"Liste scrollt nach (Doppel-)Klick von allein: {bar.value()}"

    class Wheel:      # Mausrad-Rastschritt ohne Pixel-Delta
        def __init__(self, delta): self._delta = delta
        def angleDelta(self): return QPoint(0, self._delta)
        def pixelDelta(self): return QPoint(0, 0)
        def accept(self): pass

    assert thumbs.WHEEL_STEP <= 60, "Mausrad-Scrollen der Bilder soll ruhig sein (höchstens ein halbes Bild pro Rastschritt)"
    thumbs.wheelEvent(Wheel(-120))
    assert bar.value() < thumbs.WHEEL_STEP, "Mausrad muss nachgleiten, nicht springen"
    deadline = time.time() + 3
    while thumbs._glide.state() == thumbs._glide.State.Running and time.time() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert bar.value() == thumbs.WHEEL_STEP, f"ein Rastschritt = {thumbs.WHEEL_STEP} px, gemessen {bar.value()}"
    thumbs.wheelEvent(Wheel(-120))
    thumbs.wheelEvent(Wheel(-120))              # schnelles Drehen addiert sich
    while thumbs._glide.state() == thumbs._glide.State.Running and time.time() < deadline + 3:
        app.processEvents()
        time.sleep(0.01)
    assert bar.value() == 3 * thumbs.WHEEL_STEP, f"drei Rastschritte = {3 * thumbs.WHEEL_STEP} px, gemessen {bar.value()}"
    thumbs.set_current_row_visible(19)          # Programm darf gezielt scrollen
    assert bar.value() > 3 * thumbs.WHEEL_STEP
    thumbs.close()

    # ---- Treffer der Textsuche: gleiches sanftes Scrollen, Doppelklick wählt das Lied vor ------
    results = window.results_list
    assert isinstance(results, SmoothListWidget), "Trefferliste muss sanft scrollen wie die Bilder"
    assert results.WHEEL_STEP <= 60 and hasattr(results, "_glide"), "Trefferliste braucht das sanfte Nachgleiten"

    # Seitenleiste ist so hoch wie das Fenster, auch wenn die Ansagen aufgeklappt sind (Trennlinie wächst mit)
    assert window.side_panel.height() == window.root.height(), "Seitenleiste muss die volle Fensterhöhe haben"
    heights_open = [s.height() for s in camera_panel.slots]
    window.pausentimer_panel.toggle()           # zuklappen
    settle()
    assert window.height() == window.COLLAPSED_HEIGHT
    assert window.side_panel.height() == window.root.height() == window.COLLAPSED_HEIGHT
    assert [s.height() for s in camera_panel.slots] == heights_open, "Kamerabilder dürfen beim Zuklappen nicht springen"
    total = sum(heights_open) + camera_panel.SLOT_GAP * (len(heights_open) - 1)
    assert total <= camera_panel.height(), "zwei 16:9-Bilder müssen ins eingeklappte Fenster passen"
    window.pausentimer_panel.toggle()           # wieder aufklappen
    settle()
    assert window.side_panel.height() > window.COLLAPSED_HEIGHT, "Seitenleiste wächst nicht mit den Ansagen"
    _field = window.images_panel.empty_label if window.images_panel.count() == 0 else window.images_panel.list
    _text = window.pausentimer_panel.text_edit
    assert _field.mapTo(window, _P(0, 0)).y() == _text.mapTo(window, _P(0, 0)).y(), \
        "Bilder-Feld beginnt nicht auf Höhe des Ansagen-Textfelds"
    if window.images_panel.count() == 0:
        assert _field.height() == _text.height(), "leeres Bilder-Feld ist nicht so hoch wie das Ansagen-Textfeld"
    window.side_panel.grab().save(str(SHOTS / "06_kameras.png"))

    # ---- Kamera-/Mikrofon-Anfrage beim Start ---------------------------------------------------
    from PyQt6.QtCore import Qt as _Qt

    from services import permissions as perms
    from services.settings import app_settings as _settings
    granted, denied, undetermined = (_Qt.PermissionStatus.Granted, _Qt.PermissionStatus.Denied,
                                     _Qt.PermissionStatus.Undetermined)
    saved = (perms.status, perms.request, perms._show_intro, perms._show_denied)
    states = {"camera": undetermined, "microphone": undetermined}
    calls = []

    def fake_request(permission):
        calls.append(("request", permission.key))
        states[permission.key] = granted if permission.key == "camera" else denied
        return states[permission.key]

    perms.status = lambda permission: states[permission.key]
    perms.request = fake_request
    perms._show_intro = lambda pending, parent=None: calls.append(("intro", [p.key for p in pending]))
    perms._show_denied = lambda missing, parent=None: calls.append(("denied", [p.key for p in missing]))
    try:
        perms.ensure_at_start()
        assert calls == [("intro", ["camera", "microphone"]), ("request", "camera"), ("request", "microphone"),
                         ("denied", ["microphone"])], calls
        calls.clear()
        perms.ensure_at_start()                 # nichts offen, Mikrofon weiter verweigert -> Hinweis erscheint erneut
        assert calls == [("denied", ["microphone"])], calls
        _settings().setValue("permission_hint_off_microphone", True)
        calls.clear()
        perms.ensure_at_start()                 # „nicht mehr zeigen“ ist gesetzt
        assert calls == [], calls
        states["microphone"] = granted          # wieder erlaubt -> Merker wird gelöscht
        perms.ensure_at_start()
        assert not _settings().value("permission_hint_off_microphone", False, type=bool)
    finally:
        perms.status, perms.request, perms._show_intro, perms._show_denied = saved
    # Echte Statusabfrage. Dialoge und die Systemabfrage sind ersetzt: auf einem frischen Mac (z. B. GitHub-Runner)
    # ist der Zugriff noch nie entschieden – ein echter Hinweisdialog würde den Test sonst für immer blockieren.
    shown = []
    perms._show_intro = lambda pending, parent=None: shown.append("intro")
    perms._show_denied = lambda missing, parent=None: shown.append("denied")
    perms.request = lambda permission: perms.status(permission)
    try:
        perms.ensure_at_start()
    finally:
        perms.status, perms.request, perms._show_intro, perms._show_denied = saved
    if sys.platform != "darwin":                # nur macOS fragt wirklich nach
        assert shown == [], f"auf diesem System darf kein Hinweis erscheinen: {shown}"

    # ---- Auto-Updater --------------------------------------------------------------------------
    import datetime as _dt
    import time

    from services import updater as up
    from services.recorder import recording_filename
    from services.settings import app_settings as _s
    from ui.updater_ui import KEY_CHECK_ON_START, KEY_SKIPPED, UpdateController
    assert recording_filename("Max Muster", _dt.datetime(2026, 10, 18)) == "2026.10.18 Max Muster.wav"
    assert recording_filename("", _dt.datetime(2026, 10, 18)) == "2026.10.18.wav"
    assert recording_filename('a/b:c?', _dt.datetime(2026, 10, 18)) == "2026.10.18 abc.wav"

    # Liste der gespielten Lieder und Testmodus
    from services.settings import app_settings as _ss
    _ss().setValue(_songlog.KEY_ENABLED, True)
    _ss().setValue(_songlog.KEY_TEST_MODE, False)
    found, seen = [], set()
    for query in ("1", "2", "3", "4", "5", "50", "p"):
        for song in window.db.search_songs(query):
            if song.rowid not in seen:
                seen.add(song.rowid)
                found.append(song)
    assert len(found) >= 2, "Testdatenbank braucht mindestens zwei Lieder"
    window._load_song(found[0])
    window.open_presentation()
    window.open_presentation()                                  # dasselbe Lied direkt nochmal: kein zweiter Eintrag
    log_file = _songlog.file_path()
    import datetime as _dt0
    header = _songlog.format_day(_dt0.date.today())
    expected = _songlog.format_entry(found[0].nummer, found[0].titel)
    assert log_file.read_text(encoding="utf-8").splitlines() == [header, expected], log_file.read_text(encoding="utf-8")
    window._load_song(found[1])
    window.open_presentation()
    lines = log_file.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 3 and lines[2] == _songlog.format_entry(found[1].nummer, found[1].titel)
    _ss().setValue(_songlog.KEY_TEST_MODE, True)
    window._update_title()
    assert "Testmodus" in window.windowTitle()
    window._load_song(found[0])
    window.open_presentation()
    assert len(log_file.read_text(encoding="utf-8").splitlines()) == 3, "im Testmodus darf nichts gespeichert werden"
    _ss().setValue(_songlog.KEY_ENABLED, False)
    _ss().setValue(_songlog.KEY_TEST_MODE, False)
    window._update_title()
    assert "Testmodus" not in window.windowTitle()
    window.open_presentation()
    assert len(log_file.read_text(encoding="utf-8").splitlines()) == 3, "ausgeschaltet darf nichts gespeichert werden"
    _ss().setValue(_songlog.KEY_ENABLED, True)                  # neuer Tag: Leerzeile, Kopfzeile mit Wochentag, weiter in derselben Datei
    assert _songlog.record("7", "Neuer Tag", when=_dt0.date(2026, 10, 25))
    assert not _songlog.record("7", "Neuer Tag", when=_dt0.date(2026, 10, 25)), "doppelt hintereinander"
    assert _songlog.record("8", None, when=_dt0.date(2026, 10, 25))
    assert log_file.read_text(encoding="utf-8").splitlines()[3:] == ["", "Sonntag, 25.10.2026", "7 – Neuer Tag", "8"]
    assert _songlog.format_day(_dt0.date(2026, 10, 21)) == "Mittwoch, 21.10.2026"
    # sonntags 9–10 Uhr / mittwochs 18:30–19:30 Uhr: in der Historie, aber mit „[nicht melden]“
    assert _songlog.record("9", "Kinderstunde", at=_dt0.datetime(2026, 10, 25, 9, 15))
    assert _songlog.record("10", "Bibelstunde", at=_dt0.datetime(2026, 10, 28, 19, 0))
    assert _songlog.record("11", "Nachmittag", at=_dt0.datetime(2026, 10, 28, 20, 0))
    _rows = log_file.read_text(encoding="utf-8").splitlines()
    assert "9 – Kinderstunde  [nicht melden]" in _rows and "10 – Bibelstunde  [nicht melden]" in _rows \
        and "11 – Nachmittag" in _rows, "Zeitfenster ohne CCLI-Meldung werden nicht markiert"
    _ss().setValue(_songlog.KEY_TEST_MODE, True)                # ab hier: Testmodus, damit spätere Tests nichts schreiben
    _ss().setValue(_songlog.KEY_ENABLED, True)
    window.close_beamer_windows()
    from ui.settings_dialog import SettingsDialog as _SD
    dlg = _SD(window)
    assert dlg.songlog_checkbox.isChecked() and dlg.songlog_checkbox.toolTip()
    assert not hasattr(dlg, "test_mode_checkbox") and not hasattr(dlg, "open_songlog_btn")
    dlg.close()
    assert _songlog.format_entry("p012", "None") == "p012" and _songlog.format_entry(5, " Ein  Titel ") == "5 – Ein Titel"

    # CCLI-Meldetabelle: Nummern kommen aus der CCLI-Spalte der Datenbank
    from services import ccli as _ccli
    _ccli.report_path = lambda: tmp / "CCLI-Meldung.csv"
    ccli_db = tmp / "ccli_test.db"
    make_test_db(tmp, "ccli_test.db")
    con = sqlite3.connect(ccli_db)
    con.execute("ALTER TABLE Data ADD COLUMN CCLI_Nr TEXT")
    con.execute("UPDATE Data SET CCLI_Nr = 'CCLI 7 123 456' WHERE nummer = 'p001'")
    con.commit()
    con.close()
    assert _ccli.numbers_from_db(ccli_db) == {"p001": "7123456"}, "CCLI-Nummer muss auf Ziffern bereinigt werden"
    ccli_log = tmp / "ccli_log.txt"
    log_text = "Sonntag, 25.10.2026\r\np001 – Eins\r\n7 – Sieben\r\n8\r\n"
    ccli_log.write_bytes(log_text.encode("utf-8"))      # Bytes: Windows würde \n sonst zu \r\n machen

    class _Store(dict):
        def value(self, key, default=None):
            return self.get(key, default)

        def setValue(self, key, value):
            self[key] = value

    _store = _Store()
    out, count, missing = _ccli.export_report(ccli_db, log_path=ccli_log, store=_store)
    assert out.read_text(encoding="utf-8-sig").split() == ["7123456"] and count == 1 and missing == 2, \
        "in der Meldeliste stehen nur CCLI-Nummern"
    assert ccli_log.read_bytes() == log_text.encode("utf-8"), "die Liste der gespielten Lieder (Historie) bleibt unverändert"
    out, count, missing = _ccli.export_report(ccli_db, log_path=ccli_log, store=_store)
    assert count == 1 and missing == 0, "schon übernommene Lieder dürfen nicht noch einmal in die Meldeliste"
    with open(ccli_log, "a", encoding="utf-8", newline="") as _f:
        _f.write("\r\nSonntag, 01.11.2026\r\np001 – Eins\r\n7 – Sieben\r\np001 – Eins\r\n")
    out, count, missing = _ccli.export_report(ccli_db, log_path=ccli_log, store=_store)
    assert out.read_text(encoding="utf-8-sig").split() == ["7123456", "7123456"] and count == 2, \
        "neue Nummern kommen hinter offene, dasselbe Lied am selben Tag nur einmal"
    _ccli.write_report([], out)
    assert not out.exists(), "ist alles gemeldet, wird die Meldeliste gelöscht"
    for _name in _ccli.LEGACY_FILES:
        (tmp / _name).write_text("x", encoding="utf-8")
    _ccli.remove_legacy_files(tmp)
    assert not any((tmp / _name).exists() for _name in _ccli.LEGACY_FILES), "alte Ergebnis-/Archivdateien müssen weg"
    assert _ccli.parse_log("Sonntag, 25.10.2026\n7 – A – B\n8\n") == [("Sonntag", "25.10.2026", "7", "A – B"),
                                                                       ("Sonntag", "25.10.2026", "8", "")]
    from PyQt6.QtGui import QFontMetrics as _FM
    from PyQt6.QtWidgets import QMessageBox as _MB

    from ui import theme as _theme
    _box = _MB(window)
    _long = _box.addButton("Herunterladen und installieren", _MB.ButtonRole.AcceptRole)
    _box.show()
    _theme._DialogButtonFit.fit(_box)
    assert _long.width() >= _FM(_long.font()).horizontalAdvance(_long.text()) + 2 * _theme.DIALOG_BUTTON_PAD, \
        "Knopftext im Meldungsfenster ist abgeschnitten"
    _box.close()
    # Autostart (Linux-Testlauf: Eintrag in einem Testordner statt im echten ~/.config)
    from services import autostart as _autostart
    if sys.platform.startswith("linux"):
        _old_xdg = os.environ.get("XDG_CONFIG_HOME")
        os.environ["XDG_CONFIG_HOME"] = str(tmp / "xdg")
        try:
            assert not _autostart.is_enabled()
            assert _autostart.set_enabled(True) == "" and _autostart.is_enabled(), "Autostart lässt sich nicht einschalten"
            assert "main.py" in _autostart._linux_desktop().read_text(encoding="utf-8")
            assert _autostart.set_enabled(False) == "" and not _autostart.is_enabled(), "Autostart lässt sich nicht ausschalten"
        finally:
            if _old_xdg is None:
                os.environ.pop("XDG_CONFIG_HOME", None)
            else:
                os.environ["XDG_CONFIG_HOME"] = _old_xdg
    dlg = _SD(window)
    assert dlg.autostart_checkbox.isChecked() == _autostart.is_enabled(), "Autostart-Häkchen zeigt den falschen Zustand"
    assert dlg.no_report_input.text() == _songlog.format_windows(_songlog.no_report_windows()), "Zeitfenster fehlen im Dialog"
    assert _songlog.parse_windows("So 9:00-10:00; Mittwoch 18.30–19.30 Uhr") == list(_songlog.NO_REPORT_WINDOWS)
    assert dlg.ccli_report_btn.toolTip() and not hasattr(dlg, "ccli_export_btn") and not hasattr(dlg, "ccli_numbers_btn")
    dlg.close()

    test_updater_logic(tmp)
    test_server_sync(tmp)
    test_remote_relay(tmp)
    test_remote_site()
    test_new_features(tmp)
    test_xair()
    test_history_and_problem_report(tmp)
    assert not UpdateController(window, repo="").available, "ohne Repository muss der Updater aus sein"
    dialog = SettingsDialog(window, updates_available=False)
    assert not dialog.update_check_btn.isEnabled() and "nicht eingerichtet" in dialog.update_status_label.text()
    dialog.close()
    dialog = SettingsDialog(window, updates_available=True)
    assert dialog.update_check_btn.isEnabled() and dialog.update_checkbox.isChecked(), "Update-Suche ist standardmäßig an"
    requested = []
    dialog.update_check_requested.connect(lambda: requested.append(1))
    dialog.update_check_btn.click()
    assert requested == [1]
    dialog.update_checkbox.setChecked(False)
    dialog.save_and_close()
    assert _s().value(KEY_CHECK_ON_START, True, type=bool) is False
    _s().setValue(KEY_CHECK_ON_START, True)

    newer = up.UpdateInfo("9.9.9", "Neu", "LiedManager-9.9.9-Windows-Setup.exe", "http://127.0.0.1:9/x", 1, "", "")
    controller = UpdateController(window, repo="o/r", fetch=lambda repo: {})
    offers, messages = [], []
    controller._offer = lambda info: offers.append(info.version)
    controller.status_changed.connect(messages.append)
    controller._on_checked(newer, manual=False)
    assert offers == ["9.9.9"] and "9.9.9" in messages[-1]
    _s().setValue(KEY_SKIPPED, "9.9.9")
    controller._on_checked(newer, manual=False)
    assert offers == ["9.9.9"], "übersprungene Version darf beim Start nicht erneut angeboten werden"
    controller._on_checked(newer, manual=True)
    assert offers == ["9.9.9", "9.9.9"], "manuelle Suche zeigt auch übersprungene Versionen"
    _s().setValue(KEY_SKIPPED, "")
    assert controller.blocking_reason() == "", "ohne laufende Dienste darf installiert werden"
    window.pausentimer._native._running = True
    assert "Pausentimer" in controller.blocking_reason()
    window.pausentimer._native._running = False

    # ---- Pausentimer: Folienwechsel hängen weder vom Fenster-Fokus noch von einem hängenden Skript ab ---
    import datetime

    from services import power
    from services.powerpoint import goto_command
    ctrl = window.pausentimer._powerpoint                 # PowerPoint-Pausentimer (Einstellung „PowerPoint“)
    calls = []
    ctrl._run_script = lambda command, callback, timeout_ms=0: calls.append(command)
    ctrl._running, ctrl._slide = True, None
    ctrl._finish = datetime.datetime.now() + datetime.timedelta(seconds=200)
    ctrl._check_slide()                                   # 5 Minuten vor Ende -> Folie 2, ohne dass etwas angewählt ist
    assert calls == [goto_command(2)], f"Folie 2 wurde nicht angefordert: {calls}"
    ctrl._slide, ctrl._finish = 2, datetime.datetime.now() + datetime.timedelta(seconds=59.5)
    ctrl._check_slide()                                   # letzte Minute: 60 s -> Folie 3, Takt kommt von uns
    assert calls == [goto_command(2), goto_command(3)], f"Folie 3 wurde nicht angefordert: {calls}"
    from services.powerpoint import slide_for_seconds_left as sfl
    assert [sfl(x) for x in (400, 301, 300, 61, 60, 59.5, 30.2, 1.5, 1, 0.4, 0, -5)] == \
        [1, 1, 2, 2, 3, 3, 32, 61, 62, 62, 63, 63], "Folie zur Restzeit stimmt nicht"
    del ctrl._run_script
    ctrl._running = False

    # ---- Pausenzeit: immer wieder 11:30 ----------------------------------------------------------
    from PyQt6.QtCore import QTime
    panel = window.pausentimer_panel
    assert panel.end_time() == QTime(11, 30), "die Pausenzeit muss beim Start 11:30 sein"
    panel.time_edit.setTime(QTime(10, 5))
    panel.set_running(True)
    assert panel.end_time() == QTime(10, 5), "während der Pause darf die Zeit geändert bleiben"
    panel.set_running(False)
    assert panel.end_time() == QTime(11, 30), "nach der Pause muss die Zeit wieder 11:30 sein"

    hang = []
    ctrl._run_script((sys.executable, ["-c", "import time; time.sleep(30)"]),
                     lambda ok, out: hang.append((ok, out)), timeout_ms=400)
    deadline = time.time() + 5
    while not hang and time.time() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert hang and hang[0][0] is False and "Zeitüberschreitung" in hang[0][1], f"hängendes Skript wird nicht beendet: {hang}"
    assert ctrl._script_busy is False, "nach einem hängenden Skript muss der nächste Folienwechsel wieder möglich sein"

    test_pausenfolie(tmp, window, app)

    awake = power.keep_awake()                            # Rechner/Beamer wach, Programm im Hintergrund nicht gedrosselt
    if sys.platform == "darwin" or sys.platform.startswith("win"):
        assert awake and power.is_active(), "Wach-halten (kein App Nap / Ruhezustand) funktioniert nicht"
    power.release()
    assert not power.is_active()

    results = []
    threaded = UpdateController(window, repo="o/r", fetch=lambda repo: {"tag_name": "v1.0.0", "assets": []})
    threaded._on_checked = lambda info, manual: results.append(info)   # „aktuell“ -> None
    threaded.check(manual=False)
    deadline = time.time() + 5
    while threaded._check_worker is not None and time.time() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert results == [None], f"Update-Suche im Hintergrund lieferte {results}"

    # ---- Datenbank wechseln (Einstellungen -> Daten) -------------------------------------------
    from services.settings import app_settings
    other = tmp / "andere"
    other.mkdir()
    make_test_db(other, "Zweite.db", [("p050", None, "Nun danket alle Gott", "Nun danket alle Gott\nmit Herzen, Mund und Händen")])
    app_settings().setValue("database_path", str(other / "Zweite.db"))
    window._reload_database_if_changed()
    assert window.db.db_name == str(other / "Zweite.db"), "Datenbank wurde nicht gewechselt"
    assert window.current_song_summary is None, "Liedauswahl muss nach dem Wechsel leer sein"
    assert window.db.search_songs("50") and not window.db.search_songs("1"), "falscher Datenbestand nach dem Wechsel"
    app_settings().setValue("database_path", "")
    window._reload_database_if_changed()
    assert window.db.search_songs("1"), "Zurück zur automatischen Datenbank hat nicht geklappt"

    # ---- Datenbank-Abgleich: Einstellungen und Hintergrund-Abgleich -----------------------------
    from ui.settings_dialog import SettingsDialog
    from ui.sync_ui import KEY_SYNC_TOKEN, KEY_SYNC_URL
    sync_server = FakeSyncServer()
    try:
        assert window.sync.url() == "", "Abgleich muss ohne Server-Adresse aus sein"
        dialog = SettingsDialog(window)
        assert dialog.sync_url_input.text() == "" and "Aus" in dialog.sync_status_label.text()
        assert not hasattr(dialog, "sync_checkbox") and not dialog.sync_now_btn.isEnabled()
        for widget in (dialog.sync_url_input, dialog.sync_token_input, dialog.sync_now_btn,
                       dialog.sync_force_btn):
            assert widget.toolTip(), "neue Elemente brauchen einen Hover-Hinweis"
        asked = []
        dialog.sync_requested.connect(lambda url, token: asked.append((url, token)))
        dialog.sync_url_input.setText(sync_server.url)
        dialog.sync_token_input.setText(FakeSyncServer.TOKEN)
        dialog._apply_sync_fields()                    # wirkt sofort, sobald eine Adresse da ist
        assert window.sync.enabled() and window.sync._timer.isActive()
        dialog.sync_now_btn.click()
        assert asked == [(sync_server.url, FakeSyncServer.TOKEN)], "„Jetzt abgleichen“ muss Adresse und Token aus den Feldern melden"
        assert window.sync.url() == sync_server.url and window.sync.token() == FakeSyncServer.TOKEN, \
            "Adresse und Token müssen schon beim Abgleichen gespeichert sein (auch ohne „Speichern“)"
        dialog.save_and_close()
        assert window.sync.url() == sync_server.url and window.sync.token() == FakeSyncServer.TOKEN
        window.sync.configure()
        assert window.sync._timer.isActive()
        assert window.sync.run_now(), "Abgleich muss im Hintergrund starten"
        deadline = time.time() + 15
        while window.sync._worker is not None and time.time() < deadline:   # bis der Abgleich fertig gemeldet ist
            app.processEvents()
            time.sleep(0.02)
        settle()
        assert "Eingerichtet" in window.sync.status, f"erster Abgleich: {window.sync.status}"
        assert window.db.search_songs("1"), "das Programm muss während und nach dem Abgleich benutzbar sein"
        dialog = SettingsDialog(window)
        dialog.sync_auto_checkbox.setChecked(False)    # nur noch auf Knopfdruck; Adresse und Token bleiben
        assert not window.sync._timer.isActive() and window.sync.enabled()
        assert window.sync.url() == sync_server.url and window.sync.token() == FakeSyncServer.TOKEN
        assert dialog.sync_url_input.text() == sync_server.url and dialog.sync_now_btn.isEnabled()
        dialog.close()
        app_settings().setValue(KEY_SYNC_URL, "")
        app_settings().setValue(KEY_SYNC_TOKEN, "")
        window.sync.configure()
        assert not window.sync._timer.isActive() and "Aus" in window.sync.status
    finally:
        sync_server.stop()

    # ---- Steuerung über die Website: Backend im laufenden Programm ----------------------------
    import threading

    from PyQt6.QtGui import QColor, QImage

    from services.remote_common import RemoteError as _RemoteError
    from ui.settings_dialog import SettingsDialog as _SettingsDialog
    from ui.web_remote_ui import (
        KEY_SITE_ENABLED,
        KEY_SITE_TOKEN,
        KEY_SITE_URL,
        LEGACY_KEYS,
    )

    dialog = _SettingsDialog(window)
    assert not hasattr(dialog, "remote_checkbox"), "der lokale Webserver ist entfernt"
    tab_names = [dialog.tab_group.button(i).text() for i in range(len(dialog.tab_group.buttons()))]
    assert tab_names == ["Allgemein", "Pausentimer", "Lieder", "Audio", "Kamera", "Remote"], tab_names
    page = dialog.pages.widget
    assert page(0).isAncestorOf(dialog.update_checkbox) and page(0).isAncestorOf(dialog.autostart_checkbox)
    assert page(0).isAncestorOf(dialog.screen_combo)
    assert page(1).isAncestorOf(dialog.open_presentations_btn), "Pausenfolien-Ordner gehört zur Pausentimer-Vorlage"
    assert page(2).isAncestorOf(dialog.db_path_input) and page(2).isAncestorOf(dialog.sync_url_input)
    assert page(2).isAncestorOf(dialog.songlog_checkbox) and page(2).isAncestorOf(dialog.ccli_report_btn)
    assert page(3).isAncestorOf(dialog.rec_path_input) and page(3).isAncestorOf(dialog.wp_checkbox)
    assert page(3).isAncestorOf(dialog.xair_ip_input)
    assert page(4).isAncestorOf(dialog.cam1_combo) and page(4).isAncestorOf(dialog.atem_ip_input)
    assert page(5).isAncestorOf(dialog.site_url_input) and not page(2).isAncestorOf(dialog.site_checkbox)
    for key in LEGACY_KEYS:                                   # alte Einstellungen (auch die PIN im Klartext) werden entfernt
        app_settings().setValue(key, "alt")
    window.remote.configure()
    assert not any(app_settings().contains(key) for key in LEGACY_KEYS), "alte Einstellungen des lokalen Webservers bleiben liegen"

    # Website-Steuerung: standardmäßig aus, Hover-Hinweise, verständliche Meldung bei unbrauchbaren Eingaben
    assert not dialog.site_checkbox.isChecked(), "Website-Steuerung muss standardmäßig aus sein"
    for widget in (dialog.site_checkbox, dialog.site_url_input, dialog.site_token_input):
        assert widget.toolTip(), "neue Elemente brauchen einen Hover-Hinweis"
    app_settings().setValue(KEY_SITE_ENABLED, True)
    app_settings().setValue(KEY_SITE_URL, "http://example.com/remote.php")
    app_settings().setValue(KEY_SITE_TOKEN, "x" * 20)
    window.remote.configure()
    assert "Nicht gestartet" in window.remote.site_status and "https://" in window.remote.site_status, window.remote.site_status
    app_settings().setValue(KEY_SITE_ENABLED, False)
    window.remote.configure()
    assert window.remote.site_status == "Aus" and window.remote._relay is None
    dialog = _SettingsDialog(window)                          # Schalter im Dialog wirkt sofort und trennt die Verbindung
    dialog.site_checkbox.setChecked(True)
    assert window.remote.site_enabled() and "Nicht gestartet" in dialog.site_status_label.text()
    dialog.site_checkbox.setChecked(False)
    assert not window.remote.site_enabled() and dialog.site_status_label.text() == "Aus" and window.remote._relay is None
    assert dialog.site_url_input.text() == "http://example.com/remote.php", "Adresse muss beim Ausschalten erhalten bleiben"
    dialog.close()

    web_image = tmp / "web.png"
    qimage = QImage(400, 250, QImage.Format.Format_RGB32)
    qimage.fill(QColor("#C08A53"))
    assert qimage.save(str(web_image))
    web_bytes = web_image.read_bytes()
    outcome = {}

    def web_client():
        try:
            backend = window.remote          # wird hier wie von der Website-Verbindung aus einem anderen Thread gerufen

            def action(name, body=None, expect=200):
                try:
                    data = backend.action(name, body or {})
                    status = 200
                except _RemoteError as exc:
                    data, status = exc.message, exc.status
                assert status == expect, f"{name}: erwartet {expect}, bekam {status} {data}"
                return data

            def state():
                return backend.state()

            first = state()
            for key in ("version", "beamer", "song", "audio", "recording", "pause", "images", "gallery", "cameras",
                        "atem", "sync"):
                assert key in first, f"Zustand ohne „{key}“"
            assert len(first["cameras"]) == 2 and first["recording"]["state"] == "idle"

            songs = action("search", {"text": "1"})["songs"]
            assert songs, "Suche liefert nichts"
            action("select", {"rowid": songs[0]["rowid"]})
            assert state()["song"]["nummer"] == songs[0]["nummer"]
            action("show_song")
            assert state()["beamer"] == "song"
            action("beamer_close")
            assert state()["beamer"] == "none"

            action("audio_play", {"slot": 0}, expect=409)
            action("images_open", {"index": 0}, expect=409)
            action("gallery_step", {"delta": 1}, expect=409)
            action("record_save", expect=409)
            action("camera_fullscreen", {"index": 0}, expect=409)
            action("sync_now", expect=409)
            action("atem", {"kind": "zzz"}, expect=400)
            action("select", {"rowid": "x"}, expect=400)
            action("gibt_es_nicht", expect=404)

            data = backend.upload("handy.png", web_bytes)
            assert data["added"] == 1, data
            images = state()["images"]["items"]
            assert len(images) == 1 and images[0]["name"].endswith("handy.png")
            thumb = backend.thumbnail(0)
            assert thumb and thumb.startswith(b"\xff\xd8"), "Vorschaubild ist kein JPEG"
            action("images_open", {"index": 0})
            current = state()
            assert current["beamer"] == "gallery" and current["gallery"] == {"index": 0, "count": 1}
            action("gallery_close")
            assert state()["beamer"] == "none"
            action("images_remove", {"index": 0})
            assert state()["images"]["items"] == []

            action("pause_set", {"lines": ["Hallo", "Welt"], "end": "12:15"})
            pause = state()["pause"]
            assert pause["lines"] == ["Hallo", "Welt"] and pause["end"] == "12:15" and not pause["running"], pause
            action("escape")
            outcome["ok"] = True
        except BaseException:
            import traceback as _tb
            outcome["error"] = _tb.format_exc()

    client = threading.Thread(target=web_client, daemon=True)
    client.start()
    deadline = time.time() + 90
    while client.is_alive() and time.time() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert outcome.get("ok"), f"Website-Steuerung: {outcome.get('error', 'Zeitüberschreitung')}"
    assert window.pausentimer_panel.text_edit.toPlainText() == "Hallo\nWelt", "Ansagen aus dem Browser fehlen im Programm"
    window.pausentimer_panel.text_edit.setPlainText("")
    window.close_beamer_windows()
    window.images_panel.clear()
    window.side_panel.show_tab(TAB_CAMERA)

    # ---- Bilder: Ablage, Vorschau, Galerie ----------------------------------------------------
    folder = tmp / "bilder"
    make_test_images(folder)

    panel = window.images_panel
    window.side_panel.show_tab(TAB_IMAGES)
    settle()
    assert camera_panel.isVisible() and window.side_panel.current_tab() == TAB_CAMERA, \
        "ohne Reiterleiste bleibt immer die Kamera-Seite sichtbar"
    assert not any(b.isVisible() for b in window.side_panel.group.buttons()), "die Reiterknöpfe sind entfernt"
    window.side_panel.show_tab(TAB_CAMERA)

    window.add_images([str(folder)])     # wie Drag & Drop: legt ab und wechselt selbst auf „Bilder“
    settle()
    assert panel.count() == 3, f"erwartet 3 Bilder, bekommen: {panel.count()}"
    assert window.side_panel.current_tab() == TAB_CAMERA, "ohne Reiterleiste bleibt die Kamera-Seite stehen"
    assert not hasattr(panel, "hint"), "mit Bildern gehört der ganze Platz der Vorschau (keine Erklärung darunter)"
    assert window.side_panel.group.button(TAB_IMAGES).badge == 3, "Reiter Bilder zeigt die Anzahl nicht"
    assert panel.paths()[2].endswith("bild10.png"), "natürliche Sortierung (bild2 vor bild10) fehlt"
    window.side_panel.grab().save(str(SHOTS / "10_bilder_vorschau.png"))

    again, _ = panel.add_paths([str(folder / "bild1.png")])
    assert again == 0, "Duplikat wurde nochmal aufgenommen"

    window.open_gallery(1)
    settle()
    gallery = window.gallery_window
    assert gallery is not None and gallery.isVisible(), "Galerie wurde nicht geöffnet"
    gallery.grab().save(str(SHOTS / "11_galerie_bild2.png"))
    assert panel.list.currentRow() == 1, "das Bild auf dem Beamer muss in der Vorschau gewählt (heller Rahmen) sein"
    gallery.step(1)
    assert gallery._index == 2
    assert panel.list.currentRow() == 2, "heller Rahmen muss dem Beamer-Bild folgen"
    assert gallery._fade_anim.state() == QAbstractAnimation.State.Running, "Überblendung startet nicht"
    deadline = time.monotonic() + 2.0
    while gallery._fade_old is not None and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)
    assert gallery._fade_old is None, "Überblendung endet nicht"
    settle()
    gallery.step(1)   # am Ende bleibt es beim letzten Bild und zeigt kurz den Hinweis „Ende“
    assert gallery._index == 2
    assert gallery._end_label.isVisible(), "„Ende“-Hinweis fehlt"
    gallery.grab().save(str(SHOTS / "13_galerie_ende.png"))
    gallery.step(-5)
    assert gallery._index == 0
    assert panel.list.currentRow() == 0, "heller Rahmen muss beim Zurückblättern mitwandern"
    assert not gallery._end_label.isVisible(), "„Ende“-Hinweis muss beim Zurückblättern verschwinden"
    gallery._on_mouse_activity()
    settle()
    gallery.grab().save(str(SHOTS / "12_galerie_mit_buttons.png"))

    window.handle_escape()           # Esc schließt zuerst die Galerie
    settle()
    assert window.gallery_window is None, "Esc hat die Galerie nicht geschlossen"

    # ---- Auf dem Beamer ist immer nur eines: Neues schließt, was vorher offen war --------------
    song = window.db.search_songs("1")[0]
    window.current_song_full = window.db.get_song_full_data(song.rowid)
    window.open_gallery(0)
    settle()
    assert window.gallery_window is not None
    window.open_presentation()                   # Lied ersetzt die Galerie
    settle()
    assert window.presentation_window is not None and window.gallery_window is None, \
        "Lied auf dem Beamer muss die Bilder-Galerie schließen"
    window.open_gallery(0)                       # Galerie ersetzt das Lied
    settle()
    assert window.gallery_window is not None and window.presentation_window is None, \
        "Bilder-Galerie muss das Lied auf dem Beamer schließen"
    window.open_presentation()
    window.camera_panel.fullscreen_opening.emit()  # ein Kamerabild geht auf den Beamer
    settle()
    assert window.gallery_window is None and window.presentation_window is None, \
        "Kamera-Vollbild muss Lied und Galerie schließen"
    window.pausentimer.running_changed.emit(True)
    window.open_gallery(0)
    window.pausentimer.running_changed.emit(True)  # Pausenfolie ersetzt ebenfalls alles
    settle()
    assert window.gallery_window is None, "Pausentimer muss die Galerie auf dem Beamer schließen"

    panel.remove_selected()
    assert panel.count() == 2
    panel.clear()
    settle()
    assert panel.count() == 0, "leere Ablage muss leer sein"
    assert window.side_panel.group.button(TAB_IMAGES).badge == 0, "Reiter muss nach dem Leeren ohne Zahl sein"

    print("Smoke-Test bestanden. Screenshots:", SHOTS)
    sys.stdout.flush()
    sys.stderr.flush()
    os._exit(0)   # sofort beenden: das Aufräumen der Qt-Objekte beim Verlassen von main() stürzt sonst ab


if __name__ == "__main__":
    try:
        code = main()
    except BaseException:
        traceback.print_exc()
        code = 1
    sys.stdout.flush()
    sys.stderr.flush()
    # Harter Ausstieg: verhindert, dass Qt-Aufräumarbeiten (Kamera/Audio) den Prozess in CI hängen lassen
    os._exit(code)
