"""Pausentimer: steuert eine PowerPoint-Vorlage (macOS: AppleScript, Windows: PowerShell/COM).

Alle Skript-Aufrufe laufen asynchron (QProcess), die Oberfläche bleibt dabei bedienbar.
Die Folie wechselt automatisch: Start -> 5 Minuten vor Ende -> 1 Minute vor Ende.
"""
import base64
import datetime
import logging
import math
import os
import shutil
import sys
import tempfile
import threading

from PyQt6.QtCore import QObject, QProcess, QTime, QTimer, pyqtSignal

from paths import custom_pausentimer_dir
from services import atem, power
from services.settings import app_settings

log = logging.getLogger(__name__)

SLIDE_START, SLIDE_5MIN, SLIDE_1MIN = 1, 2, 3
SLIDE_END = 63                 # „Ende!“; die Folien 3 bis 62 zeigen 60 bis 1 (eine Folie pro Sekunde)

SCRIPT_TIMEOUT_MS = 20_000     # Folienwechsel/Schließen: länger darf ein Skript nicht hängen
OPEN_TIMEOUT_MS = 90_000       # PowerPoint starten (große Datei, erste Rückfrage zur Automatisierung am Mac)


def slide_for_seconds_left(seconds_left: float) -> int:
    """Folie zur Restzeit. In der letzten Minute zählt die Vorlage sekundenweise herunter (Folie 3 = „60“ …
    Folie 62 = „1“, Folie 63 = „Ende!“). Den Takt gibt LiedManager vor, nicht PowerPoint: dessen eigene
    Zeitschaltung der Folien bleibt stehen, sobald PowerPoint nicht das aktive Programm ist."""
    if seconds_left <= 0:
        return SLIDE_END
    if seconds_left <= 60:
        return min(SLIDE_END - 1, max(SLIDE_1MIN, SLIDE_END - math.ceil(seconds_left)))
    if seconds_left <= 300:
        return SLIDE_5MIN
    return SLIDE_START


def template_path(season: str):
    """Die Vorlage „pausentimer_<Jahreszeit>.pptx“ im Datenordner (…/LiedManager/pausentimer). Mitgeliefert wird keine."""
    return custom_pausentimer_dir() / f"pausentimer_{season}.pptx"


def write_announcements(pptx_path: str, lines) -> None:
    """Schreibt die Ansagen in die erste Folie und behält dabei Schrift, Größe, Farbe und Stil der Vorlage."""
    try:
        from pptx import Presentation
    except ImportError:
        log.info("python-pptx nicht installiert – Ansagen werden nicht in die Folie geschrieben.")
        return

    prs = Presentation(pptx_path)
    if len(prs.slides) == 0:
        return
    shapes = list(prs.slides[0].shapes)
    target = (next((s for s in shapes if s.name == "AnsagenTextfeld"), None)
              or next((s for s in shapes if s.has_text_frame), None))
    if target is None or not target.has_text_frame:
        return

    tf = target.text_frame
    paragraphs = list(tf.paragraphs)

    sample = None  # (name, size, color, bold, italic) der ersten Textzeile
    if paragraphs:
        source = paragraphs[0].runs[0].font if paragraphs[0].runs else paragraphs[0].font
        sample = (source.name, source.size, source.color,
                  getattr(source, "bold", None), getattr(source, "italic", None))

    for i, line in enumerate(lines):
        paragraph = paragraphs[i] if i < len(paragraphs) else tf.add_paragraph()
        paragraph.text = line
        if sample is None:
            continue
        name, size, color, bold, italic = sample
        for run in paragraph.runs:
            if name:
                run.font.name = name
            if size:
                run.font.size = size
            if color:
                try:
                    run.font.color.rgb = color.rgb
                except Exception:
                    pass
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic

    for paragraph in paragraphs[len(lines):]:
        paragraph.text = ""

    prs.save(pptx_path)


# Es soll nur die Präsentation zu sehen sein – ohne Referentenansicht und ohne Bearbeitungsfenster.
# macOS: Das Bearbeitungsfenster lässt sich bei PowerPoint nicht abschalten. Es wird deshalb klein hinter die laufende
# Bildschirmpräsentation geschoben (dort ist es verdeckt und erscheint nicht im Dock). Klappt das nicht, wird es wie
# bisher minimiert. Die Befehle laufen als Text-Skript in „try“: kennt eine PowerPoint-Version einen Begriff nicht, wird
# er übergangen, statt dass der Pausentimer nicht startet.
def _applescript_literal(text: str) -> str:
    """Mehrzeiliges AppleScript als Zeichenketten-Ausdruck (für „run script“)."""
    lines = [line.replace("\\", "\\\\").replace('"', '\\"') for line in text.strip().splitlines()]
    return '"' + '" & linefeed & "'.join(lines) + '"'


_HIDE_DOCS_INNER = """
tell application "Microsoft PowerPoint"
    set moved to false
    try
        set b to bounds of slide show window 1
        set x to item 1 of b
        set y to item 2 of b
        repeat with w in (get every document window)
            set bounds of w to {x + 40, y + 40, x + 640, y + 440}
        end repeat
        set moved to true
    end try
    if not moved then
        repeat with w in (get every window)
            try
                set window state of w to minimized
            end try
        end repeat
    end if
end tell
"""
_MINIMIZE_ALL = f"""
    try
        run script {_applescript_literal(_HIDE_DOCS_INNER)}
    end try
""".replace("{", "{{").replace("}", "}}")     # die Skripte darunter werden mit .format() gefüllt

_OPEN_SCRIPT = r'''
tell application "Microsoft PowerPoint"
    open POSIX file "{path}"
    try
        run script "tell application \"Microsoft PowerPoint\" to set show presenter view of slide show settings of active presentation to false"
    end try
    run slide show slide show settings of active presentation
''' + _MINIMIZE_ALL + r'''
    delay 0.8
''' + _MINIMIZE_ALL + r'''
end tell
'''

_CLOSE_SCRIPT = '''
tell application "Microsoft PowerPoint"
    if (count of presentations) > 0 then
        close active presentation saving no
    end if
    quit
end tell
'''

_GOTO_SCRIPT = '''
tell application "Microsoft PowerPoint"
    if (count of slide show windows) is 0 then return "wait"
    set theView to slide show view of slide show window 1
''' + _MINIMIZE_ALL + r'''    try
        go to slide theView number {slide}
        return "ok"
    on error errMsg1
        try
            go to first slide theView
            repeat {steps} times
                go to next slide theView
            end repeat
            return "ok"
        on error errMsg2
            return "error: " & errMsg1 & " | " & errMsg2
        end try
    end try
end tell
'''


# --- Windows: PowerShell + PowerPoint-COM (ohne zusätzliche Pakete) ---
# Präsentation ohne Bearbeitungsfenster öffnen (Presentations.Open … WithWindow = 0). Das Skript meldet danach die
# Fenster-Kennungen („HWND:<PowerPoint>:<Präsentation>“); das Ausblenden aus Taskleiste und Alt+Tab erledigt Python
# (hide_windows_from_taskbar) – schnell und ohne in PowerShell C#-Code zu übersetzen (das dauerte Sekunden und hing teils).
_PS_HWND = ("try {{ Write-Output ('HWND:' + $ppt.HWND + ':' + $ppt.SlideShowWindows.Item(1).HWND) }} catch {{}}")
_PS_OPEN = """
$ErrorActionPreference = 'Stop'
try {{
  $ppt = New-Object -ComObject PowerPoint.Application
  $pres = $ppt.Presentations.Open('{path}', 0, 0, 0)
  try {{ $pres.SlideShowSettings.ShowPresenterView = 0 }} catch {{}}
  $pres.SlideShowSettings.Run() | Out-Null
  {hwnd}
  exit 0
}} catch {{ Write-Output $_.Exception.Message; exit 1 }}
"""

_PS_CLOSE = """
try {
  $ppt = [Runtime.InteropServices.Marshal]::GetActiveObject('PowerPoint.Application')
  foreach ($p in @($ppt.Presentations)) { $p.Saved = -1; $p.Close() }
  $ppt.Quit()
} catch {}
exit 0
"""

_PS_GOTO = """
$ErrorActionPreference = 'Stop'
try {{
  $ppt = [Runtime.InteropServices.Marshal]::GetActiveObject('PowerPoint.Application')
  if ($ppt.SlideShowWindows.Count -eq 0) {{ Write-Output 'wait'; exit 0 }}
  $ppt.SlideShowWindows.Item(1).View.GotoSlide({slide})
  Write-Output 'ok'
}} catch {{ Write-Output ('error: ' + $_.Exception.Message); exit 1 }}
"""


def _hide_console(process: QProcess) -> None:
    """Windows: kein kurz aufblitzendes Konsolenfenster beim Start von PowerShell (auf dem Mac ohne Wirkung)."""
    if not sys.platform.startswith("win"):
        return
    def no_window(arguments):
        arguments.flags |= 0x08000000  # CREATE_NO_WINDOW
    try:
        process.setCreateProcessArgumentsModifier(no_window)
    except (AttributeError, TypeError):  # nur unter Windows vorhanden
        pass


def hide_windows_from_taskbar(output: str) -> None:
    """Windows: blendet das leere PowerPoint-Hauptfenster aus und nimmt das Präsentationsfenster aus Taskleiste und Alt+Tab
    (es bleibt auf dem Beamer sichtbar). ``output`` enthält „HWND:<PowerPoint>:<Präsentation>“ vom Öffnen-Skript."""
    if not sys.platform.startswith("win"):
        return
    line = next((row for row in str(output).splitlines() if row.startswith("HWND:")), "")
    parts = line.split(":")
    if len(parts) != 3:
        return
    try:
        import ctypes
        user32 = ctypes.windll.user32
        app, show = int(parts[1] or 0), int(parts[2] or 0)
        GWL_EXSTYLE, WS_EX_TOOLWINDOW, WS_EX_APPWINDOW = -20, 0x80, 0x40000
        SW_HIDE, SW_SHOWNA = 0, 8
        if app and user32.IsWindowVisible(app):
            user32.ShowWindow(app, SW_HIDE)
        if show:
            style = user32.GetWindowLongW(show, GWL_EXSTYLE)
            wanted = (style | WS_EX_TOOLWINDOW) & ~WS_EX_APPWINDOW
            if wanted != style:
                user32.ShowWindow(show, SW_HIDE)          # die Taskleiste übernimmt den neuen Stil erst nach Aus-/Einblenden
                user32.SetWindowLongW(show, GWL_EXSTYLE, wanted)
                user32.ShowWindow(show, SW_SHOWNA)
                user32.SetForegroundWindow(show)          # Präsentation bleibt aktiv (sonst halten die Folienzeiten an)
    except (AttributeError, OSError, ValueError) as exc:
        log.info("Präsentationsfenster nicht aus der Taskleiste genommen: %s", exc)


def clear_download_mark(path: str) -> None:
    """Windows: Die Markierung „aus dem Internet geladen“ (Zone.Identifier) der Kopie entfernen. Sonst öffnet PowerPoint
    die Datei in der geschützten Ansicht – die Bildschirmpräsentation startet dann nicht oder hängt."""
    if sys.platform.startswith("win"):
        try:
            os.remove(path + ":Zone.Identifier")
        except OSError:
            pass


def supported() -> bool:
    return sys.platform == "darwin" or sys.platform.startswith("win")


class _Command(tuple):
    """(Programm, Argumente); unter Windows zusätzlich ``session_script``: dasselbe als Einzeiler für die dauerhafte PowerShell."""
    session_script = ""


def _command(mac_script: str, ps_script: str, session_script: str = ""):
    """(Programm, Argumente) für das jeweilige Betriebssystem."""
    if sys.platform == "darwin":
        return _Command(("osascript", ["-e", mac_script]))
    encoded = base64.b64encode(ps_script.encode("utf-16-le")).decode("ascii")
    command = _Command(("powershell.exe", ["-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                                           "-EncodedCommand", encoded]))
    command.session_script = session_script
    return command


# Dieselben Schritte als Einzeiler für die dauerhaft laufende PowerShell (ohne „exit“, Fehler werden abgefangen).
_SESSION_OPEN = ("$ErrorActionPreference = 'Stop'; $ppt = New-Object -ComObject PowerPoint.Application; "
                 "$pres = $ppt.Presentations.Open('{path}', 0, 0, 0); "
                 "try {{ $pres.SlideShowSettings.ShowPresenterView = 0 }} catch {{}}; "
                 "$pres.SlideShowSettings.Run() | Out-Null; {hwnd}")
_SESSION_CLOSE = ("try {{ if ($null -eq $ppt) {{ $ppt = [Runtime.InteropServices.Marshal]::GetActiveObject('PowerPoint.Application') }}; "
                  "foreach ($p in @($ppt.Presentations)) {{ $p.Saved = -1; $p.Close() }}; $ppt.Quit() }} catch {{}}; $ppt = $null")
_SESSION_GOTO = ("$ErrorActionPreference = 'Stop'; "
                 "try {{ $lmshows = $ppt.SlideShowWindows.Count }} catch {{ "      # $ppt aus dem Öffnen weiterverwenden
                 "$ppt = [Runtime.InteropServices.Marshal]::GetActiveObject('PowerPoint.Application'); "
                 "$lmshows = $ppt.SlideShowWindows.Count }}; "
                 "if ($lmshows -eq 0) {{ Write-Output 'wait' }} else "
                 "{{ $ppt.SlideShowWindows.Item(1).View.GotoSlide({slide}); Write-Output 'ok' }}")


def open_command(path: str):
    ps_path = path.replace("'", "''")
    hwnd = _PS_HWND.format()
    return _command(_OPEN_SCRIPT.format(path=path), _PS_OPEN.format(path=ps_path, hwnd=hwnd),
                    _SESSION_OPEN.format(path=ps_path, hwnd=hwnd))


def close_command():
    return _command(_CLOSE_SCRIPT, _PS_CLOSE, _SESSION_CLOSE.format())


def goto_command(slide: int):
    return _command(_GOTO_SCRIPT.format(slide=slide, steps=slide - 1), _PS_GOTO.format(slide=slide),
                    _SESSION_GOTO.format(slide=slide))


SESSION_MARKER = "__LM_DONE__"


class _PowerShellSession(QObject):
    """Eine dauerhaft laufende PowerShell (Windows): spart den Start von ``powershell.exe`` bei jedem Folienwechsel (0,5–1 s).

    Jeder Befehl ist ein Einzeiler; danach schreibt die Sitzung eine Marke mit dem Ergebnis (0 = gut). Die Antwort wird bis zur
    Marke gesammelt. Stirbt der Prozess oder reagiert er nicht, räumt der Aufrufer ihn ab; der nächste Befehl startet neu.
    """
    def __init__(self, parent=None):
        super().__init__(parent)
        self._process = None
        self._buffer = ""
        self._callback = None

    def ensure(self) -> bool:
        process = self._process
        if process is not None and process.state() != QProcess.ProcessState.NotRunning:
            return True
        process = QProcess(self)
        process.setProcessChannelMode(QProcess.ProcessChannelMode.MergedChannels)
        _hide_console(process)
        process.readyReadStandardOutput.connect(self._on_output)
        process.finished.connect(self._on_died)
        process.errorOccurred.connect(lambda err: self._on_died() if err == QProcess.ProcessError.FailedToStart else None)
        process.start("powershell.exe", ["-NoProfile", "-NoLogo", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                                         "-Command", "-"])
        if not process.waitForStarted(5000):
            process.deleteLater()
            return False
        self._process = process
        self._buffer = ""
        return True

    def submit(self, one_liner: str, callback) -> None:
        self._buffer = ""
        self._callback = callback
        line = (f"$lmcode = 0; try {{ {one_liner} }} catch {{ Write-Output $_.Exception.Message; $lmcode = 1 }}; "
                f"Write-Output \"{SESSION_MARKER}$lmcode\"\n")
        self._process.write(line.encode("utf-8"))

    def _on_output(self) -> None:
        if self._process is None:
            return
        self._buffer += bytes(self._process.readAllStandardOutput()).decode("utf-8", "replace")
        if SESSION_MARKER not in self._buffer:
            return
        text, _, rest = self._buffer.partition(SESSION_MARKER)
        code = rest.strip().splitlines()[0].strip() if rest.strip() else "1"
        self._buffer = ""
        callback, self._callback = self._callback, None
        if callback is not None:
            callback(code == "0", text.strip())

    def _on_died(self, *_args) -> None:
        callback, self._callback = self._callback, None
        if callback is not None:
            callback(False, "PowerShell wurde beendet")

    def kill(self) -> None:
        process, self._process = self._process, None
        self._callback = None
        if process is not None:
            try:
                process.finished.disconnect()
            except TypeError:
                pass
            process.kill()
            process.deleteLater()


class PausentimerController(QObject):
    """Startet/stoppt die Pausenpräsentation und wechselt die Folien passend zur Restzeit.

    Signale:
        running_changed(bool)   True, sobald der Start angefordert wurde; False nach Stopp oder Fehler
        error(str)              Meldung für den Benutzer
    """
    running_changed = pyqtSignal(bool)
    error = pyqtSignal(str)
    _prepared = pyqtSignal(str, str)     # (Pfad der vorbereiteten Datei, Fehlertext) – vom Vorbereitungs-Thread

    def __init__(self, parent=None):
        super().__init__(parent)
        self._running = False
        self._preparing = False          # Vorlage wird gerade im Hintergrund vorbereitet
        self._start_cancelled = False
        self._start_args = ()
        self._end_time = QTime(11, 30)
        self._finish = None
        self._slide = None
        self._script_busy = False
        self._last_problem = ""      # zuletzt protokollierte Störung beim Folienwechsel (nicht jede Sekunde loggen)

        self._timer = QTimer(self)
        self._timer.setInterval(250)   # nur zum Prüfen; ein Skript läuft erst, wenn sich die Folie ändern muss
        self._timer.timeout.connect(self._check_slide)
        self._prepared.connect(self._on_prepared)
        self._session = None             # dauerhafte PowerShell (Windows)
        self._session_failures = 0       # drei Fehler in Folge: zurück zum Einzelstart je Befehl

    @property
    def running(self) -> bool:
        return self._running

    def shutdown(self) -> None:
        """Beim Beenden des Programms: die vorab gestartete PowerShell beenden (PowerPoint selbst bleibt unberührt)."""
        self._timer.stop()
        if self._session is not None:
            self._session.kill()

    def warm_up(self) -> None:
        """Windows: die PowerShell schon im Leerlauf starten, damit der Pausentimer auf Knopfdruck schneller öffnet."""
        if not sys.platform.startswith("win") or self._running:
            return
        if self._session is None:
            self._session = _PowerShellSession(self)
        self._session.ensure()

    def seconds_left(self):
        """Restzeit der Pause in Sekunden (``None``, wenn der Pausentimer nicht läuft)."""
        if not self._running or self._finish is None:
            return None
        return max(0, int((self._finish - datetime.datetime.now()).total_seconds()))

    # ------------------------------------------------------------------
    # Öffentliche Steuerung
    # ------------------------------------------------------------------
    def start(self, lines, end_time: QTime) -> None:
        if self._running or self._preparing:
            return
        if not supported():
            self.error.emit("Der Pausentimer wird nur unter macOS und Windows unterstützt.")
            return
        season = app_settings().value("presentation_season", "frühling", type=str)
        source = template_path(season)
        if not source.exists():
            self.error.emit(f"Keine Pausenfolie gefunden: „{source.name}“.\n\nLege deine PowerPoint-Vorlage unter diesem Namen in "
                            f"den Ordner\n{source.parent}\n(Einstellungen → Anzeige → „Ordner für Pausenfolien öffnen“).")
            return

        self._preparing = True
        self._start_args = (end_time,)
        lines = list(lines)

        def prepare():                           # Kopieren und Ansagen eintragen (kann bei großen Vorlagen Sekunden dauern)
            try:
                temp_path = os.path.join(tempfile.gettempdir(), f"pausentimer_temp_{season}.pptx")
                shutil.copyfile(source, temp_path)       # nur der Inhalt (keine Markierung „aus dem Internet“)
                clear_download_mark(temp_path)
                try:
                    write_announcements(temp_path, lines)
                except Exception as exc:  # fehlerhafte Ansagen sollen den Start nicht verhindern
                    log.warning("Ansagen konnten nicht in die Vorlage geschrieben werden: %s", exc)
                self._prepared.emit(temp_path, "")
            except OSError as exc:
                self._prepared.emit("", str(exc))

        threading.Thread(target=prepare, name="pausentimer-prepare", daemon=True).start()

    def _on_prepared(self, temp_path: str, problem: str) -> None:
        self._preparing = False
        if self._start_cancelled:
            self._start_cancelled = False
            return
        if problem:
            self.error.emit(f"Die Vorlage konnte nicht vorbereitet werden:\n{problem}")
            return
        (end_time,) = self._start_args
        atem.switch("presentation")
        self._end_time = end_time
        self._finish = self._compute_finish()
        self._slide = None
        self._running = True
        power.keep_awake()    # Rechner/Beamer nicht einschlafen lassen, Programm im Hintergrund nicht drosseln
        threading.Thread(target=power.exempt_from_app_nap, args=(power.POWERPOINT_BUNDLE_ID,),
                         name="app-nap", daemon=True).start()     # „defaults write“ nicht im Hauptthread
        self.running_changed.emit(True)

        self._run_script(open_command(os.path.abspath(temp_path)), self._on_open_finished,
                         timeout_ms=OPEN_TIMEOUT_MS)

    def stop(self) -> None:
        if self._preparing:
            self._start_cancelled = True          # Start läuft noch im Hintergrund: danach nichts mehr öffnen
            return
        if not self._running:
            return
        self._timer.stop()
        self._running = False
        self._slide = None
        power.release()
        self._run_script(close_command(), lambda ok, out: None)   # Sitzung bleibt offen: der nächste Start geht schneller
        self.running_changed.emit(False)
        atem.switch("default")

    def set_end_time(self, end_time: QTime) -> None:
        """Endzeit wurde geändert -> Restzeit neu berechnen, Folie neu bestimmen."""
        self._end_time = end_time
        if self._running:
            self._finish = self._compute_finish()
            self._slide = None

    # ------------------------------------------------------------------
    # Intern
    # ------------------------------------------------------------------
    def _compute_finish(self) -> datetime.datetime:
        now = datetime.datetime.now()
        finish = datetime.datetime.combine(
            now.date(), datetime.time(self._end_time.hour(), self._end_time.minute())
        )
        if finish <= now:
            finish += datetime.timedelta(days=1)
        return finish

    def _on_open_finished(self, ok: bool, output: str) -> None:
        if not self._running:
            return
        if ok:
            hide_windows_from_taskbar(output)
            self._timer.start()
            return
        log.error("PowerPoint-Präsentation konnte nicht gestartet werden: %s", output)
        self._running = False
        power.release()
        self.running_changed.emit(False)
        atem.switch("default")
        self.error.emit("Die PowerPoint-Präsentation konnte nicht gestartet werden.\n"
                        "Ist Microsoft PowerPoint installiert?")

    def _check_slide(self) -> None:
        if not self._running or self._script_busy:
            return
        seconds_left = (self._finish - datetime.datetime.now()).total_seconds()
        target = slide_for_seconds_left(seconds_left)
        if target == self._slide:
            return
        def done(ok: bool, output: str, target=target):
            if ok and output == "ok":
                self._slide = target
                self._last_problem = ""
                return
            problem = output or "keine Antwort von PowerPoint"
            if problem != self._last_problem:   # Folienwechsel wird laufend neu versucht
                self._last_problem = problem
                log.warning("Folienwechsel auf Folie %d nicht möglich: %s", target, problem)

        self._run_script(goto_command(target), done)

    def _run_in_session(self, script: str, callback, timeout_ms: int) -> bool:
        """Windows: Befehl in der dauerhaften PowerShell ausführen. ``False``: Sitzung nicht verfügbar, bitte klassisch starten."""
        if self._session is None:
            self._session = _PowerShellSession(self)
        if not self._session.ensure():
            self._session_failures += 1
            return False
        state = {"called": False}
        self._script_busy = True
        watchdog = QTimer(self)
        watchdog.setSingleShot(True)
        watchdog.setInterval(timeout_ms)

        def finish(ok: bool, output: str) -> None:
            if state["called"]:
                return
            state["called"] = True
            watchdog.stop()
            watchdog.deleteLater()
            self._script_busy = False
            self._session_failures = 0 if ok or output not in ("PowerShell wurde beendet", "Zeitüberschreitung") else \
                self._session_failures + 1
            callback(ok, output)

        def on_timeout() -> None:
            log.warning("PowerShell-Sitzung reagiert nicht (%.1f s) – wird neu gestartet.", timeout_ms / 1000)
            self._session.kill()
            finish(False, "Zeitüberschreitung")

        watchdog.timeout.connect(on_timeout)
        watchdog.start()
        self._session.submit(script, finish)
        return True

    def _run_script(self, command, callback, timeout_ms: int = SCRIPT_TIMEOUT_MS) -> None:
        """Führt (Programm, Argumente) asynchron aus; ``callback(ok, stdout)`` wird genau einmal aufgerufen.

        Hängt das Skript länger als ``timeout_ms`` (z. B. PowerPoint antwortet nicht), wird es beendet und als
        Fehler gemeldet. Sonst bliebe ``_script_busy`` für immer gesetzt und es gäbe keine Folienwechsel mehr.
        """
        script = getattr(command, "session_script", "")
        if script and sys.platform.startswith("win") and self._session_failures < 3 and self._run_in_session(
                script, callback, timeout_ms):
            return
        program, args = command
        process = QProcess(self)
        _hide_console(process)
        state = {"called": False}
        self._script_busy = True

        watchdog = QTimer(process)
        watchdog.setSingleShot(True)
        watchdog.setInterval(timeout_ms)

        def release():
            try:
                process.deleteLater()
            except RuntimeError:  # Objekt ist schon zerstört (z. B. beim Beenden des Programms)
                pass

        def finish(ok: bool, output: str):
            if state["called"]:
                return
            state["called"] = True
            watchdog.stop()
            self._script_busy = False
            # Ein per kill() beendeter Prozess läuft kurz weiter: erst löschen, wenn er wirklich beendet ist
            # (siehe on_finished), sonst meldet Qt „Destroyed while process is still running“.
            try:
                idle = process.state() == QProcess.ProcessState.NotRunning
            except RuntimeError:
                idle = False
            if idle:
                release()
            callback(ok, output)

        def on_timeout():
            if state["called"]:
                return
            log.warning("Skript für den Pausentimer reagiert nicht (%.1f s) – wird beendet.", timeout_ms / 1000)
            process.kill()
            finish(False, "Zeitüberschreitung")

        def on_finished(code, _status):
            if state["called"]:  # schon per Zeitüberschreitung abgeschlossen: nur noch aufräumen
                release()
                return
            try:
                output = bytes(process.readAllStandardOutput()).decode("utf-8", "replace").strip()
            except RuntimeError:  # Objekt wird gerade zerstört (z. B. beim Beenden des Programms)
                output = ""
                code = -1
            finish(code == 0, output)

        def on_error(err):
            if err == QProcess.ProcessError.FailedToStart:
                finish(False, f"{err}")

        watchdog.timeout.connect(on_timeout)
        watchdog.start()
        process.finished.connect(on_finished)
        process.errorOccurred.connect(on_error)
        process.start(program, args)
